from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK, WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Smart_Study_Planner_Project_Guide.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color="DADCE0", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, width_dxa=9360, indent_dxa=120):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")


def add_code_block(doc, code, language="text"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_width(table)
    set_table_borders(table, color="DADCE0", size="4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F6F8FA")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    paragraph = cell.paragraphs[0]
    paragraph.style = doc.styles["CodeBlock"]
    run = paragraph.add_run(sanitize_xml_text(code.strip("\n")))
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(36, 41, 47)
    caption = doc.add_paragraph()
    caption.style = doc.styles["CaptionSmall"]
    caption.add_run("Code/command: " + language)


def sanitize_xml_text(text):
    result = []
    for ch in text:
        if ch == "\x08":
            result.append("\\b")
        elif ch in ("\n", "\t", "\r") or ord(ch) >= 32:
            result.append(ch)
    return "".join(result)


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_numbered(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.add_run(item)


def add_note(doc, text, title="Ghi nho"):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table)
    set_table_borders(table, color="BFC7D5", size="4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, "EEF3FA")
    paragraph = cell.paragraphs[0]
    paragraph.style = doc.styles["Normal"]
    title_run = paragraph.add_run(title + ": ")
    title_run.bold = True
    paragraph.add_run(text)


def add_error_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_width(table)
    set_table_borders(table)
    headers = ["Loi gap phai", "Nguyen nhan", "Cach fix"]
    widths = [2600, 3200, 3560]
    for idx, header in enumerate(headers):
        cell = table.cell(0, idx)
        set_cell_shading(cell, "E8EEF5")
        set_cell_width(cell, widths[idx])
        paragraph = cell.paragraphs[0]
        paragraph.style = doc.styles["TableText"]
        run = paragraph.add_run(header)
        run.bold = True

    for issue, cause, fix in rows:
        cells = table.add_row().cells
        for idx, text in enumerate((issue, cause, fix)):
            set_cell_width(cells[idx], widths[idx])
            paragraph = cells[idx].paragraphs[0]
            paragraph.style = doc.styles["TableText"]
            paragraph.add_run(text)


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(11, 37, 69)
    title.paragraph_format.space_after = Pt(10)

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    code = styles.add_style("CodeBlock", 1)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    code.font.size = Pt(8.5)
    code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1.0

    caption = styles.add_style("CaptionSmall", 1)
    caption.font.name = "Calibri"
    caption.font.size = Pt(8)
    caption.font.color.rgb = RGBColor(85, 85, 85)
    caption.paragraph_format.space_after = Pt(6)

    table_text = styles.add_style("TableText", 1)
    table_text.font.name = "Calibri"
    table_text.font.size = Pt(9)
    table_text.paragraph_format.space_after = Pt(0)
    table_text.paragraph_format.line_spacing = 1.1

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
    footer.style = styles["Normal"]
    footer.add_run("Smart Study Planner - Project learning guide")

    return doc


def build_doc():
    doc = setup_document()

    doc.add_paragraph("Smart Study Planner", style="Title")
    subtitle = doc.add_paragraph()
    subtitle.add_run("Tai lieu hoc lai du an C++ + SQL Server + FastAPI").bold = True
    doc.add_paragraph(
        "Muc tieu cua file nay la ghi lai qua trinh lam du an theo tung buoc: "
        "lam gi, vi sao lam, gap loi gi, sua nhu the nao, va code trong moi phan co tac dung gi."
    )

    doc.add_heading("1. Tong quan du an", level=1)
    doc.add_paragraph(
        "Du an bat dau tu nen tang ban da hoc: database o truong, mot so thuat toan co ban C++, "
        "va mong muon co mot san pham ca nhan vua phai de dua vao profile. Y tuong duoc chon la "
        "mot ung dung lap ke hoach hoc tap bang flashcard, co SQL Server luu du lieu va C++ xu ly logic."
    )
    add_bullets(
        doc,
        [
            "Ten du an: Smart Study Planner.",
            "Muc tieu nguoi dung: them cau hoi, tim kiem, on lai flashcard den han, xem chu de yeu.",
            "Muc tieu ky thuat: thuc hanh SQL schema, C++ ODBC, CMake, Git/GitHub, va chuan bi len web.",
            "Kien truc hien tai: main.cpp -> Menu -> Service -> Repository -> DbConnection -> SQL Server.",
        ],
    )
    add_code_block(
        doc,
        """
main.cpp
-> Menu
-> Service
-> Repository
-> DbConnection
-> SQL Server
        """,
        "architecture",
    )
    doc.add_paragraph(
        "Y nghia cua kien truc nay: main.cpp chi dieu huong chuong trinh; Menu chi hien thi lua chon; "
        "Service xu ly nghiep vu va input; Repository chua SQL; DbConnection phu trach ket noi ODBC. "
        "Nho tach nhu vay, khi chuyen len web ta co the giu y tuong Service/Repository va chi thay lop giao dien."
    )

    doc.add_heading("2. Chuan bi cong cu", level=1)
    doc.add_paragraph("Nhung cong cu da dung trong qua trinh lam du an:")
    add_bullets(
        doc,
        [
            "Visual Studio Code: mo project, sua file C++/Python/HTML/CSS/JS, mo terminal.",
            "Visual Studio 2026 / MSVC: cung cap compiler cl.exe de build C++ tren Windows.",
            "CMake: sinh build folder va goi compiler de build file .exe.",
            "SQL Server Express 2026 va SSMS 22: tao database, chay schema, kiem tra bang du lieu.",
            "ODBC Driver 18 for SQL Server: cau noi de C++ va Python ket noi SQL Server.",
            "Git va GitHub: luu lich su code, commit, push len remote.",
            "Python 3.12, FastAPI, Uvicorn: chuan bi backend web.",
        ],
    )
    add_note(
        doc,
        "Trong du an nay, SQL Server la trung tam du lieu. C++ CLI va FastAPI backend deu ket noi vao cung database SmartStudyPlanner.",
    )

    doc.add_heading("3. Tao va ket noi SQL Server", level=1)
    doc.add_paragraph(
        "Buoc dau la dam bao SQL Server Express dang chay va SSMS ket noi duoc vao instance HOANE\\SQLEXPRESS. "
        "Trong man hinh Connect cua SSMS, dung Windows Authentication, tick Trust Server Certificate neu Encrypt la Mandatory."
    )
    add_code_block(
        doc,
        """
Server Name: HOANE\\SQLEXPRESS
Authentication: Windows Authentication
Encrypt: Mandatory
Trust Server Certificate: checked
Database Name: <default> hoac master khi tao database
        """,
        "SSMS connection",
    )
    doc.add_paragraph(
        "Tai sao ban dau co luc chon master? Vi database SmartStudyPlanner chua ton tai thi phai ket noi vao master truoc, "
        "sau do chay script tao database. Khi database da tao xong, co the chon SmartStudyPlanner trong dropdown."
    )
    add_code_block(
        doc,
        """
IF DB_ID(N'SmartStudyPlanner') IS NULL
BEGIN
    CREATE DATABASE SmartStudyPlanner;
END;
GO

USE SmartStudyPlanner;
GO
        """,
        "sql",
    )
    doc.add_paragraph(
        "Doan SQL tren kiem tra database da ton tai chua. Neu chua co thi tao moi. Lenh USE chuyen ngu canh query sang database cua du an."
    )

    doc.add_heading("4. Cau hinh SQL Server Configuration Manager", level=1)
    doc.add_paragraph(
        "Trong qua trinh lam, co luc tim SQL Server Configuration Manager nhung Start Menu chi hien search web. "
        "Cach dung dung la mo SQL Server Configuration Manager, bam SQL Server Services de xem service, roi mo SQL Server Network Configuration -> Protocols for SQLEXPRESS."
    )
    add_numbered(
        doc,
        [
            "Mo SQL Server Configuration Manager.",
            "Neu ben phai hien 'There are no items to show', chon dung node SQL Server Services hoac refresh.",
            "Dam bao SQL Server (SQLEXPRESS) dang Running.",
            "Vao SQL Server Network Configuration -> Protocols for SQLEXPRESS.",
            "Enable TCP/IP neu can ket noi qua network/local port.",
            "Restart service SQL Server (SQLEXPRESS) de ap dung thay doi.",
        ],
    )
    doc.add_paragraph(
        "Tac dung: SQL Server service la may chu database. Protocols quy dinh cach app noi vao server. "
        "TCP/IP giup client nhu SSMS, C++ ODBC, Python pyodbc ket noi on dinh hon."
    )

    doc.add_heading("5. Thiet ke database schema", level=1)
    doc.add_paragraph(
        "Database gom 4 bang chinh. Quan he la Subject co nhieu Topic, Topic co nhieu Flashcard, Flashcard co nhieu ReviewLog."
    )
    add_code_block(
        doc,
        """
Subjects 1---N Topics 1---N Flashcards 1---N ReviewLogs
        """,
        "relationship",
    )
    add_code_block(
        doc,
        """
CREATE TABLE dbo.Subjects (
    SubjectId INT IDENTITY(1,1) PRIMARY KEY,
    Name NVARCHAR(100) NOT NULL UNIQUE,
    CreatedAt DATETIME2 NOT NULL DEFAULT SYSUTCDATETIME()
);

CREATE TABLE dbo.Topics (
    TopicId INT IDENTITY(1,1) PRIMARY KEY,
    SubjectId INT NOT NULL,
    Name NVARCHAR(100) NOT NULL,
    Priority TINYINT NOT NULL DEFAULT 3,
    CONSTRAINT FK_Topics_Subjects
        FOREIGN KEY (SubjectId) REFERENCES dbo.Subjects(SubjectId)
);

CREATE TABLE dbo.Flashcards (
    FlashcardId INT IDENTITY(1,1) PRIMARY KEY,
    TopicId INT NOT NULL,
    Question NVARCHAR(500) NOT NULL,
    Answer NVARCHAR(1000) NOT NULL,
    Difficulty TINYINT NOT NULL DEFAULT 3,
    NextReviewAt DATETIME2 NOT NULL DEFAULT SYSUTCDATETIME(),
    ReviewIntervalDays INT NOT NULL DEFAULT 1,
    CONSTRAINT FK_Flashcards_Topics
        FOREIGN KEY (TopicId) REFERENCES dbo.Topics(TopicId)
);

CREATE TABLE dbo.ReviewLogs (
    ReviewLogId INT IDENTITY(1,1) PRIMARY KEY,
    FlashcardId INT NOT NULL,
    ReviewedAt DATETIME2 NOT NULL DEFAULT SYSUTCDATETIME(),
    WasCorrect BIT NOT NULL,
    CONSTRAINT FK_ReviewLogs_Flashcards
        FOREIGN KEY (FlashcardId) REFERENCES dbo.Flashcards(FlashcardId)
);
        """,
        "sql",
    )
    doc.add_paragraph(
        "Giai thich: IDENTITY tu tang ID. NVARCHAR luu duoc Unicode. FOREIGN KEY giu quan he giua bang. "
        "NextReviewAt va ReviewIntervalDays phuc vu thuat toan on tap gian cach."
    )

    doc.add_heading("6. Tao project C++ va CMake", level=1)
    doc.add_paragraph(
        "Sau khi co database, can tao ung dung C++ co the build thanh file .exe. CMakeLists.txt khai bao chuan C++20, cac file source, thu muc include va thu vien odbc32."
    )
    add_code_block(
        doc,
        """
cmake_minimum_required(VERSION 3.20)

project(SmartStudyPlanner LANGUAGES CXX)

set(CMAKE_CXX_STANDARD 20)
set(CMAKE_CXX_STANDARD_REQUIRED ON)
set(CMAKE_CXX_EXTENSIONS OFF)

add_executable(smart_study_planner
    src/DbConnection.cpp
    src/FlashcardRepository.cpp
    src/FlashcardService.cpp
    src/ReviewRepository.cpp
    src/ReviewService.cpp
    src/StatisticsRepository.cpp
    src/StatisticsService.cpp
    src/TopicRepository.cpp
    src/TopicService.cpp
    src/Utils.cpp
    src/Menu.cpp
    src/main.cpp
)

target_include_directories(smart_study_planner PRIVATE include)
target_link_libraries(smart_study_planner PRIVATE odbc32)
        """,
        "cmake",
    )
    doc.add_paragraph(
        "Giai thich: add_executable gom tat ca file .cpp can bien dich. Neu tao file .cpp moi ma quen them vao day, build se loi linker "
        "vi ham da khai bao nhung khong co phan dinh nghia duoc bien dich. odbc32 la thu vien Windows can de goi ODBC API."
    )
    add_code_block(
        doc,
        """
cmake -S . -B build -G "NMake Makefiles"
cmake --build build
.\build\smart_study_planner.exe
        """,
        "powershell",
    )

    doc.add_heading("7. Ket noi C++ voi SQL Server bang ODBC", level=1)
    doc.add_paragraph(
        "Lop DbConnection boc lai cac ham ODBC thap cap. Ben ngoai chi can goi connect, query, execute thay vi lap lai SQLAllocHandle, SQLExecDirectW o moi noi."
    )
    add_code_block(
        doc,
        """
class DbConnection {
public:
    void connect(const std::wstring& connectionString);
    void disconnect();
    bool isConnected() const;

    std::vector<std::vector<std::wstring>> query(const std::wstring& sql);
    void execute(const std::wstring& sql);

private:
    SQLHENV env_;
    SQLHDBC dbc_;
    bool connected_;
};
        """,
        "cpp",
    )
    add_code_block(
        doc,
        """
std::wstring connectionString =
    L"Driver={" + driver + L"};"
    L"Server=" + server + L";"
    L"Database=" + database + L";"
    L"Encrypt=no;"
    L"TrustServerCertificate=yes;";

if (user.empty()) {
    connectionString += L"Trusted_Connection=yes;";
} else {
    connectionString += L"UID=" + user + L";PWD=" + password + L";";
}
        """,
        "cpp",
    )
    doc.add_paragraph(
        "Giai thich: Driver chon ODBC Driver 18. Server la HOANE\\SQLEXPRESS. Database la SmartStudyPlanner. "
        "Trusted_Connection=yes dung tai khoan Windows hien tai de dang nhap SQL Server. Encrypt=no va TrustServerCertificate=yes giup tranh loi certificate trong moi truong local."
    )
    add_code_block(
        doc,
        """
auto rows = db.query(L"SELECT COUNT(*) FROM dbo.Flashcards;");
if (!rows.empty() && !rows.front().empty()) {
    std::wcout << L"Flashcards in database: " << rows.front().front() << L"\\n";
}
        """,
        "cpp",
    )
    doc.add_paragraph(
        "Doan nay la test ket noi dau tien. Neu app in duoc so flashcard, nghia la C++ -> ODBC -> SQL Server -> database da thong."
    )

    doc.add_heading("8. Tao menu CLI", level=1)
    doc.add_paragraph(
        "Ban dau menu chi co placeholder 'coming soon'. Sau do tung chuc nang duoc noi vao service rieng."
    )
    add_code_block(
        doc,
        """
int showMenu() {
    std::wcout << L"\\nSmart Study Planner\\n";
    std::wcout << L"1. List flashcards\\n";
    std::wcout << L"2. Search flashcards\\n";
    std::wcout << L"3. Add flashcard\\n";
    std::wcout << L"4. Edit flashcard\\n";
    std::wcout << L"5. Delete flashcard\\n";
    std::wcout << L"6. Review due flashcards\\n";
    std::wcout << L"7. Show weak topics\\n";
    std::wcout << L"8. Show study statistics\\n";
    std::wcout << L"9. Exit\\n";
    std::wcout << L"Choose: ";

    int choice = 0;
    std::wcin >> choice;
    std::wcin.ignore(std::numeric_limits<std::streamsize>::max(), L'\\n');
    return choice;
}
        """,
        "cpp",
    )
    doc.add_paragraph(
        "Giai thich: showMenu chi hien thi menu va tra ve lua chon. No khong nen tu query SQL, vi nhu vay menu se bi tron voi logic database."
    )
    add_code_block(
        doc,
        """
switch (choice) {
    case 1: FlashcardService::listFlashcards(db); break;
    case 2: FlashcardService::searchFlashcards(db); break;
    case 3: FlashcardService::addFlashcard(db); break;
    case 4: FlashcardService::editFlashcard(db); break;
    case 5: FlashcardService::deleteFlashcard(db); break;
    case 6: ReviewService::reviewDueFlashcards(db); break;
    case 7: ReviewService::showWeakTopics(db); break;
    case 8: StatisticsService::showStudyStatistics(db); break;
    case 9: running = false; break;
}
        """,
        "cpp",
    )

    doc.add_heading("9. Tach Model, Service, Repository", level=1)
    doc.add_paragraph(
        "Khi du an lon hon, neu de tat ca SQL trong main.cpp thi rat kho sua va kho chuyen len web. Vi vay code duoc tach thanh model, service, repository."
    )
    add_bullets(
        doc,
        [
            "Model: struct du lieu nhu Flashcard, Topic, DueReview, StudyStatistics.",
            "Service: nhan input tu user, validate, goi repository, in output.",
            "Repository: viet SQL SELECT/INSERT/UPDATE/DELETE.",
            "DbConnection: lop thap nhat de gui SQL den SQL Server.",
        ],
    )
    add_code_block(
        doc,
        """
struct Flashcard {
    int id;
    std::wstring subjectName;
    std::wstring topicName;
    std::wstring question;
    std::wstring answer;
    int difficulty;
    std::wstring nextReviewAt;
};
        """,
        "cpp",
    )
    doc.add_paragraph(
        "Model giup repository tra ve du lieu co ten truong ro rang thay vi vector row[0], row[1] nam khap noi trong code."
    )

    doc.add_heading("10. List va Search flashcards", level=1)
    doc.add_paragraph(
        "Tinh nang list lay tat ca flashcard, join qua Topics va Subjects de in duoc ten mon hoc va ten chu de. "
        "Search dung LIKE de tim trong Question, Answer, TopicName, SubjectName."
    )
    add_code_block(
        doc,
        """
std::vector<Flashcard> FlashcardRepository::search(DbConnection& db, const std::wstring& keyword) {
    std::wstring pattern = L"N'%" + escapeSql(keyword) + L"%'";

    return mapFlashcards(db.query(
        L"SELECT f.FlashcardId, s.Name, t.Name, f.Question, f.Answer, "
        L"f.Difficulty, CONVERT(NVARCHAR(19), f.NextReviewAt, 120) "
        L"FROM dbo.Flashcards f "
        L"JOIN dbo.Topics t ON f.TopicId = t.TopicId "
        L"JOIN dbo.Subjects s ON t.SubjectId = s.SubjectId "
        L"WHERE f.Question LIKE " + pattern + L" "
        L"OR f.Answer LIKE " + pattern + L" "
        L"OR t.Name LIKE " + pattern + L" "
        L"OR s.Name LIKE " + pattern + L" "
        L"ORDER BY f.FlashcardId;"
    ));
}
        """,
        "cpp",
    )
    doc.add_paragraph(
        "escapeSql dung de nhan doi dau nhay don. Neu user nhap cau hoi co dau ', SQL string se khong bi vo."
    )
    add_code_block(
        doc,
        """
std::wstring escapeSql(const std::wstring& value) {
    std::wstring escaped;
    for (wchar_t ch : value) {
        escaped.push_back(ch);
        if (ch == L'\\'') {
            escaped.push_back(L'\\'');
        }
    }
    return escaped;
}
        """,
        "cpp",
    )

    doc.add_heading("11. Add flashcard va validate Topic ID", level=1)
    doc.add_paragraph(
        "Loi quan trong da sua: khi user nhap TopicId khong ton tai, database se bao loi foreign key hoac app co the dung lai. "
        "Cach fix la kiem tra topic ton tai truoc khi insert."
    )
    add_code_block(
        doc,
        """
int topicId = readInt(L"\\nTopic ID: ", 0);
if (!TopicRepository::exists(db, topicId)) {
    std::wcout << L"Topic not found.\\n";
    return;
}
        """,
        "cpp",
    )
    add_code_block(
        doc,
        """
bool TopicRepository::exists(DbConnection& db, int topicId) {
    auto rows = db.query(
        L"SELECT COUNT(*) FROM dbo.Topics WHERE TopicId = " +
        std::to_wstring(topicId) + L";"
    );

    return !rows.empty() && !rows[0].empty() && std::stoi(rows[0][0]) > 0;
}
        """,
        "cpp",
    )
    doc.add_paragraph(
        "Giai thich: Service hoi nguoi dung nhap TopicId, Repository hoi database xem ID do co ton tai khong. "
        "Neu khong ton tai thi in Topic not found va quay lai menu, khong de database nem loi."
    )

    doc.add_heading("12. Edit va Delete flashcard", level=1)
    doc.add_paragraph(
        "Edit va Delete cung duoc dua xuong Repository. Service chi hoi input va confirm. Repository moi la noi chiu trach nhiem UPDATE/DELETE SQL."
    )
    add_code_block(
        doc,
        """
void FlashcardRepository::update(DbConnection& db, int flashcardId,
                                 const std::wstring& question,
                                 const std::wstring& answer,
                                 int difficulty) {
    std::wstring sql =
        L"UPDATE dbo.Flashcards SET "
        L"Question = N'" + escapeSql(question) + L"', "
        L"Answer = N'" + escapeSql(answer) + L"', "
        L"Difficulty = " + std::to_wstring(difficulty) + L" "
        L"WHERE FlashcardId = " + std::to_wstring(flashcardId) + L";";

    db.execute(sql);
}
        """,
        "cpp",
    )
    add_code_block(
        doc,
        """
void FlashcardRepository::remove(DbConnection& db, int flashcardId) {
    db.execute(L"DELETE FROM dbo.ReviewLogs WHERE FlashcardId = " +
               std::to_wstring(flashcardId) + L";");
    db.execute(L"DELETE FROM dbo.Flashcards WHERE FlashcardId = " +
               std::to_wstring(flashcardId) + L";");
}
        """,
        "cpp",
    )
    doc.add_paragraph(
        "Tai sao xoa ReviewLogs truoc? Vi ReviewLogs co foreign key tro den Flashcards. Neu xoa Flashcards truoc, SQL Server se chan vi van con review log tham chieu den flashcard do."
    )

    doc.add_heading("13. Review due flashcards va spaced repetition", level=1)
    doc.add_paragraph(
        "Tinh nang review lay cac flashcard den han theo NextReviewAt <= SYSUTCDATETIME(). Neu user tra loi dung, khoang on tap tang gap doi toi da 30 ngay. Neu sai, reset ve 1 ngay."
    )
    add_code_block(
        doc,
        """
SELECT TOP 10
    f.FlashcardId,
    s.Name AS SubjectName,
    t.Name AS TopicName,
    f.Question,
    f.Answer
FROM dbo.Flashcards f
JOIN dbo.Topics t ON f.TopicId = t.TopicId
JOIN dbo.Subjects s ON t.SubjectId = s.SubjectId
WHERE f.NextReviewAt <= SYSUTCDATETIME()
ORDER BY f.Difficulty DESC, f.NextReviewAt ASC;
        """,
        "sql",
    )
    add_code_block(
        doc,
        """
UPDATE dbo.Flashcards
SET ReviewIntervalDays = CASE
    WHEN correct = 1 THEN
        CASE WHEN ReviewIntervalDays * 2 > 30 THEN 30 ELSE ReviewIntervalDays * 2 END
    ELSE 1
END,
NextReviewAt = DATEADD(day, CASE
    WHEN correct = 1 THEN
        CASE WHEN ReviewIntervalDays * 2 > 30 THEN 30 ELSE ReviewIntervalDays * 2 END
    ELSE 1
END, SYSUTCDATETIME())
WHERE FlashcardId = id;
        """,
        "sql pseudo-code",
    )
    doc.add_paragraph(
        "Y nghia thuat toan: cau nao dung lien tuc se xuat hien thua dan, cau nao sai se quay lai som. "
        "Day la phien ban don gian cua spaced repetition."
    )

    doc.add_heading("14. Weak topics va statistics", level=1)
    doc.add_paragraph(
        "Weak topics dung aggregate query tren ReviewLogs de tinh ty le sai theo tung topic. Statistics gom tong flashcard, tong review, dung, sai, accuracy va so card den han."
    )
    add_code_block(
        doc,
        """
SELECT TOP 10
    s.Name AS SubjectName,
    t.Name AS TopicName,
    COUNT(rl.ReviewLogId) AS ReviewCount,
    SUM(CASE WHEN rl.WasCorrect = 0 THEN 1 ELSE 0 END) AS WrongCount,
    CAST(100.0 * SUM(CASE WHEN rl.WasCorrect = 0 THEN 1 ELSE 0 END)
         / COUNT(rl.ReviewLogId) AS DECIMAL(5,2)) AS WrongRate
FROM dbo.ReviewLogs rl
JOIN dbo.Flashcards f ON rl.FlashcardId = f.FlashcardId
JOIN dbo.Topics t ON f.TopicId = t.TopicId
JOIN dbo.Subjects s ON t.SubjectId = s.SubjectId
GROUP BY s.Name, t.Name
ORDER BY WrongRate DESC, ReviewCount DESC;
        """,
        "sql",
    )
    add_code_block(
        doc,
        """
stats.totalFlashcards = readIntResult(db, L"SELECT COUNT(*) FROM dbo.Flashcards;");
stats.totalReviews = readIntResult(db, L"SELECT COUNT(*) FROM dbo.ReviewLogs;");
stats.correctCount = readIntResult(db, L"SELECT COUNT(*) FROM dbo.ReviewLogs WHERE WasCorrect = 1;");
stats.wrongCount = readIntResult(db, L"SELECT COUNT(*) FROM dbo.ReviewLogs WHERE WasCorrect = 0;");
stats.dueTodayCount = readIntResult(db, L"SELECT COUNT(*) FROM dbo.Flashcards WHERE NextReviewAt <= SYSUTCDATETIME();");
        """,
        "cpp",
    )

    doc.add_heading("15. Git va GitHub workflow", level=1)
    doc.add_paragraph(
        "Git duoc dung de luu tung moc tien do. Sau khi build on, code duoc commit va push len GitHub."
    )
    add_code_block(
        doc,
        """
git init
git status
git add .
git commit -m "Initial C++ SQL Server study planner"
git remote add origin https://github.com/nguyenhoe11/smart-study-planner-cpp.git
git branch -M main
git push -u origin main
        """,
        "git",
    )
    doc.add_paragraph(
        "Giai thich: git init tao repo local. git add dua file vao staging. git commit tao moc lich su. "
        "git remote add origin noi repo local voi GitHub. git push day code len GitHub."
    )
    add_note(
        doc,
        "Canh bao LF/CRLF khi commit tren Windows khong phai loi nghiem trong. Do la Git thong bao se chuyen line ending khi xu ly file.",
    )

    doc.add_heading("16. Cac loi da gap va cach fix", level=1)
    add_error_table(
        doc,
        [
            (
                "Khong biet chon master hay database nao trong SSMS.",
                "Database SmartStudyPlanner chua duoc tao nen dropdown chua co.",
                "Ket noi vao master truoc, chay sql/schema.sql de tao database, sau do chon SmartStudyPlanner.",
            ),
            (
                "SQL Server Configuration Manager mo ra nhung khung ben phai khong co item.",
                "Dang chon nham node hoac chua refresh view.",
                "Chon SQL Server Services, bam refresh, sau do vao Protocols for SQLEXPRESS de cau hinh network.",
            ),
            (
                "C++ build loi khi dung cmake --build build.",
                "Terminal chua nap bien moi truong MSVC hoac CMakeLists thieu file .cpp moi.",
                "Dung Visual Studio Developer PowerShell, hoac VS Code terminal da set MSVC env; them tat ca .cpp moi vao CMakeLists.txt.",
            ),
            (
                "Ctrl+S trong VS Code lai hien loi do trong file .cpp.",
                "Thuong la IntelliSense chua bat dung include/compiler env, khong nhat thiet la build that su loi.",
                "Kiem tra bang cmake --build build. Neu build xanh thi uu tien tin build. Neu can, reload VS Code/C++ extension.",
            ),
            (
                "Nhap TopicId sai khi add flashcard lam app co nguy co loi.",
                "Flashcards.TopicId la foreign key nen ID khong ton tai bi SQL Server chan.",
                "Them TopicRepository::exists va validate truoc khi INSERT.",
            ),
            (
                "Xoa flashcard bi ket do foreign key.",
                "ReviewLogs van tham chieu FlashcardId do.",
                "Xoa ReviewLogs truoc, sau do moi xoa Flashcards.",
            ),
            (
                "VS Code bao Failed to install update khi tat app.",
                "Day la loi update cua VS Code, co the do process dang chay hoac antivirus.",
                "Khong anh huong code. Dong bot VS Code process, chay lai update sau hoac bo qua neu project van build duoc.",
            ),
            (
                "Python command bi tro sang Microsoft Store alias.",
                "WindowsApps alias chen truoc python.org trong PATH.",
                "Dung duong dan Python that hoac virtual environment: web/backend/.venv/Scripts/python.exe.",
            ),
        ],
    )

    doc.add_heading("17. Chuan bi nang cap len web", level=1)
    doc.add_paragraph(
        "Sau khi CLI da tach Service/Repository, du an duoc scaffold them web. Backend dung FastAPI de tao API, frontend dung HTML/CSS/JavaScript de goi API."
    )
    add_code_block(
        doc,
        """
web/
  backend/
    app/
      db.py
      repositories.py
      schemas.py
      main.py
    requirements.txt
    .env
  frontend/
    index.html
    styles.css
    app.js
        """,
        "project structure",
    )
    add_code_block(
        doc,
        """
def build_connection_string() -> str:
    driver = os.getenv("DB_DRIVER", "ODBC Driver 18 for SQL Server")
    server = os.getenv("DB_SERVER", r"HOANE\\SQLEXPRESS")
    database = os.getenv("DB_NAME", "SmartStudyPlanner")

    return (
        f"Driver={{{driver}}};"
        f"Server={server};"
        f"Database={database};"
        "Trusted_Connection=yes;"
        "Encrypt=no;"
        "TrustServerCertificate=yes;"
    )
        """,
        "python",
    )
    doc.add_paragraph(
        "db.py lam vai tro giong DbConnection.cpp nhung cho Python. No tao connection string va mo ket noi pyodbc den SQL Server."
    )
    add_code_block(
        doc,
        """
@app.get("/api/flashcards")
def list_flashcards():
    return repositories.list_flashcards()

@app.post("/api/flashcards", status_code=201)
def create_flashcard(payload: FlashcardCreate):
    if not repositories.topic_exists(payload.topicId):
        raise HTTPException(status_code=404, detail="Topic not found.")

    repositories.create_flashcard(
        payload.topicId,
        payload.question,
        payload.answer,
        payload.difficulty,
    )
    return {"message": "Flashcard created."}
        """,
        "python",
    )
    doc.add_paragraph(
        "main.py tao endpoint HTTP. Thay vi nguoi dung go phim trong terminal, trinh duyet se gui request GET/POST/PUT/DELETE den API."
    )
    add_code_block(
        doc,
        """
async function requestJson(url, options) {
  const response = await fetch(url, {
    headers: { "Content-Type": "application/json" },
    ...options,
  });

  if (!response.ok) {
    const error = await response.json().catch(() => ({ detail: "Request failed." }));
    throw new Error(error.detail || "Request failed.");
  }

  return response.json();
}
        """,
        "javascript",
    )
    doc.add_paragraph(
        "app.js la cau noi frontend -> backend. fetch goi API, nhan JSON, sau do render len HTML."
    )
    add_code_block(
        doc,
        """
cd C:\\Users\\Admin\\Documents\\project\\web\\backend
.\\.venv\\Scripts\\activate
uvicorn app.main:app --reload

# Mo trinh duyet:
http://127.0.0.1:8000
        """,
        "powershell",
    )

    doc.add_heading("18. Hien tai du an dang o dau", level=1)
    add_bullets(
        doc,
        [
            "CLI C++ da co day du luong chuc nang chinh.",
            "SQL Server schema da co bang can thiet, chua can doi schema.",
            "Code da tach ro Service/Repository de de nang cap.",
            "FastAPI backend da co endpoint co ban cho flashcards, topics, reviews, statistics.",
            "Frontend prototype da co dashboard, list, search, add, edit, delete, review, weak topics.",
            "Can polish UI va test ky hon truoc khi dua len portfolio.",
        ],
    )

    doc.add_heading("19. Checklist test lai du an", level=1)
    add_code_block(
        doc,
        """
cmake --build build
.\build\smart_study_planner.exe
        """,
        "powershell",
    )
    add_bullets(
        doc,
        [
            "Chon 1 de list flashcards.",
            "Chon 2 de search theo cau hoi, cau tra loi, topic hoac subject.",
            "Chon 3, nhap TopicId hop le de add.",
            "Chon 3, nhap TopicId sai de kiem tra thong bao Topic not found.",
            "Chon 4 de edit flashcard.",
            "Chon 5 de delete flashcard va confirm y.",
            "Chon 6 de review due flashcards.",
            "Chon 7 de xem weak topics.",
            "Chon 8 de xem study statistics.",
        ],
    )

    doc.add_heading("20. Lo trinh tiep theo", level=1)
    add_numbered(
        doc,
        [
            "Hoan thien UI web: bo cuc dashboard, form add/edit ro rang, review flow dep hon.",
            "Them validate frontend: difficulty 1-5, khong cho question/answer rong.",
            "Them trang topics hoac subject management neu muon quan ly mon hoc ngay tren web.",
            "Them screenshot vao README va viet mo ta portfolio bang tieng Anh.",
            "Lam demo video ngan: tao flashcard, review, xem statistics.",
            "Sau khi on dinh moi tinh den deploy hoac chuyen sang framework frontend lon hon.",
        ],
    )

    doc.add_heading("21. Tom tat bai hoc rut ra", level=1)
    add_bullets(
        doc,
        [
            "Database nen duoc thiet ke truoc vi no quy dinh du lieu ma app can xu ly.",
            "Ket noi app voi SQL Server can dung driver, server name, authentication va encryption phu hop.",
            "CMakeLists.txt phai duoc cap nhat moi khi them file .cpp moi.",
            "Service va Repository giup code de doc, de test, de chuyen len web.",
            "Foreign key bao ve du lieu nhung bat buoc app phai validate truoc khi insert/delete.",
            "Git commit theo tung moc giup quay lai duoc khi loi va giup profile trong chuyen nghiep hon.",
            "CLI tot la nen tang de lam web: khi logic da tach gon, web chi la giao dien/API moi boc ben ngoai.",
        ],
    )

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Phu luc: cac file nen doc lai khi reset chat", level=1)
    add_bullets(
        doc,
        [
            "README.md: tom tat du an, build, run, tech stack.",
            "docs/HANDOFF.md: trang thai moi nhat de tiep tuc sau khi reset chat.",
            "docs/API_CONTRACT.md: danh sach API cho web.",
            "docs/WEB_MIGRATION_PLAN.md: lo trinh chuyen CLI sang web.",
            "sql/schema.sql va sql/seed.sql: database schema va data mau.",
            "src/main.cpp, src/Menu.cpp: luong chay CLI.",
            "src/*Repository.cpp: tat ca SQL query.",
            "src/*Service.cpp: logic nghiep vu va input/output.",
            "web/backend/app/main.py: API FastAPI.",
            "web/frontend/app.js: frontend goi API bang fetch.",
        ],
    )

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_doc()
    print(OUTPUT)
