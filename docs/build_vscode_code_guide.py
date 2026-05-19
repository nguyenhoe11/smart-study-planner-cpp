from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "VSCode_Codebase_Guide.docx"


FILES_TO_INCLUDE = [
    ("CMakeLists.txt", "CMake build configuration"),
    ("include/models/Flashcard.h", "Flashcard model"),
    ("include/models/Topic.h", "Topic model"),
    ("include/models/Review.h", "Review models"),
    ("include/models/StudyStatistics.h", "Study statistics model"),
    ("include/DbConnection.h", "DbConnection declaration"),
    ("src/DbConnection.cpp", "ODBC implementation"),
    ("include/Utils.h", "Utils declaration"),
    ("src/Utils.cpp", "Input and SQL string helpers"),
    ("include/Menu.h", "Menu declaration"),
    ("src/Menu.cpp", "CLI menu"),
    ("src/main.cpp", "Program entry point"),
    ("include/FlashcardService.h", "Flashcard service declaration"),
    ("src/FlashcardService.cpp", "Flashcard user workflow"),
    ("include/FlashcardRepository.h", "Flashcard repository declaration"),
    ("src/FlashcardRepository.cpp", "Flashcard SQL queries"),
    ("include/TopicService.h", "Topic service declaration"),
    ("src/TopicService.cpp", "Topic display workflow"),
    ("include/TopicRepository.h", "Topic repository declaration"),
    ("src/TopicRepository.cpp", "Topic SQL queries"),
    ("include/ReviewService.h", "Review service declaration"),
    ("src/ReviewService.cpp", "Review user workflow"),
    ("include/ReviewRepository.h", "Review repository declaration"),
    ("src/ReviewRepository.cpp", "Review SQL queries"),
    ("include/StatisticsService.h", "Statistics service declaration"),
    ("src/StatisticsService.cpp", "Statistics display workflow"),
    ("include/StatisticsRepository.h", "Statistics repository declaration"),
    ("src/StatisticsRepository.cpp", "Statistics SQL queries"),
    ("web/backend/app/db.py", "FastAPI database helper"),
    ("web/backend/app/schemas.py", "FastAPI request schemas"),
    ("web/backend/app/repositories.py", "FastAPI SQL functions"),
    ("web/backend/app/main.py", "FastAPI routes"),
    ("web/frontend/index.html", "Static web layout"),
    ("web/frontend/app.js", "Frontend API calls and rendering"),
]


FUNCTIONS = [
    ("src/main.cpp", "wmain", "Diem bat dau cua chuong trinh C++ tren Windows. Ham nay doc tham so dong lenh, tao connection string, ket noi SQL Server, in thong tin ban dau, chay vong lap menu va goi cac Service tuong ung voi lua chon cua user."),
    ("src/Menu.cpp", "showMenu", "In danh sach lua chon trong terminal, doc so user nhap vao, xoa ky tu Enter con thua trong buffer, roi tra ve choice cho main.cpp xu ly."),
    ("src/Utils.cpp", "readLine", "In prompt ra man hinh va doc mot dong text bang std::getline. Dung khi can nhap question, answer, keyword hoac confirm y/n."),
    ("src/Utils.cpp", "readInt", "In prompt va doc so nguyen. Neu user nhap sai kieu, ham reset stream va tra ve gia tri mac dinh de app khong crash."),
    ("src/Utils.cpp", "escapeSql", "Duyet tung ky tu trong chuoi va nhan doi dau nhay don. Muc dich la tranh lam vo SQL string khi user nhap text co ky tu '."),
    ("src/DbConnection.cpp", "narrow", "Ham phu trong namespace an danh, chuyen wstring sang string ASCII don gian de dua thong bao loi ODBC vao std::runtime_error."),
    ("src/DbConnection.cpp", "DbConnection::DbConnection", "Constructor cap phat ODBC environment handle, set ODBC version, roi cap phat connection handle. Day la buoc chuan bi truoc khi connect vao SQL Server."),
    ("src/DbConnection.cpp", "DbConnection::~DbConnection", "Destructor tu dong disconnect va free cac ODBC handle. Muc dich la tranh ro ri tai nguyen khi chuong trinh ket thuc."),
    ("src/DbConnection.cpp", "DbConnection::connect", "Nhan connection string, goi SQLDriverConnectW de mo ket noi SQL Server, neu thanh cong thi danh dau connected_ = true."),
    ("src/DbConnection.cpp", "DbConnection::disconnect", "Neu dang ket noi thi goi SQLDisconnect va danh dau connected_ = false."),
    ("src/DbConnection.cpp", "DbConnection::isConnected", "Tra ve trang thai connected_. Ham nay giup cac phan khac biet DbConnection co dang mo ket noi hay khong."),
    ("src/DbConnection.cpp", "DbConnection::query", "Chay cau SELECT, doc so cot, fetch tung dong, lay du lieu tung cot bang SQLGetData va tra ve vector 2 chieu cac wstring."),
    ("src/DbConnection.cpp", "DbConnection::execute", "Chay cau SQL khong can tra ve du lieu nhu INSERT, UPDATE, DELETE. Neu ODBC bao loi thi nem exception."),
    ("src/DbConnection.cpp", "DbConnection::throwIfFailed", "Kiem tra SQLRETURN. Neu loi, doc diagnostic record tu ODBC de tao thong bao loi ro hon."),
    ("src/FlashcardService.cpp", "printFlashcards", "Ham phu de in danh sach Flashcard theo format de doc: id, subject, topic, question, difficulty va next review."),
    ("src/FlashcardService.cpp", "normalizeDifficulty", "Gioi han difficulty trong khoang 1 den 5. Neu user nhap nho hon 1 thi lay 1, lon hon 5 thi lay 5."),
    ("src/FlashcardService.cpp", "FlashcardService::listFlashcards", "Goi Repository lay tat ca flashcard roi in ra man hinh. Service khong tu viet SQL."),
    ("src/FlashcardService.cpp", "FlashcardService::searchFlashcards", "Hoi keyword, validate keyword khong rong, goi FlashcardRepository::search, roi in ket qua."),
    ("src/FlashcardService.cpp", "FlashcardService::addFlashcard", "In danh sach topic, hoi Topic ID, validate topic ton tai, hoi question/answer/difficulty, roi goi Repository insert flashcard."),
    ("src/FlashcardService.cpp", "FlashcardService::editFlashcard", "In danh sach card, hoi ID can sua, kiem tra ID hop le va ton tai, hoi noi dung moi, roi goi Repository update."),
    ("src/FlashcardService.cpp", "FlashcardService::deleteFlashcard", "In danh sach card, hoi ID can xoa, kiem tra ton tai, hoi confirm y/n, roi goi Repository xoa ReviewLogs va Flashcards."),
    ("src/FlashcardRepository.cpp", "mapFlashcards", "Ham phu chuyen ket qua query dang vector row/cell thanh vector<Flashcard>. Nho vay Service lam viec voi object co field ro rang."),
    ("src/FlashcardRepository.cpp", "FlashcardRepository::getAll", "Chay SELECT join Flashcards, Topics, Subjects de lay danh sach flashcard day du thong tin subject/topic."),
    ("src/FlashcardRepository.cpp", "FlashcardRepository::search", "Tao pattern LIKE N'%keyword%' va tim trong Question, Answer, TopicName, SubjectName."),
    ("src/FlashcardRepository.cpp", "FlashcardRepository::exists", "Kiem tra FlashcardId co ton tai khong bang SELECT COUNT(*). Dung truoc edit/delete/review."),
    ("src/FlashcardRepository.cpp", "FlashcardRepository::create", "Tao cau INSERT INTO dbo.Flashcards voi TopicId, Question, Answer, Difficulty."),
    ("src/FlashcardRepository.cpp", "FlashcardRepository::update", "Tao cau UPDATE dbo.Flashcards de sua Question, Answer va Difficulty theo FlashcardId."),
    ("src/FlashcardRepository.cpp", "FlashcardRepository::remove", "Xoa ReviewLogs truoc roi moi xoa Flashcards de khong vi pham foreign key."),
    ("src/TopicService.cpp", "TopicService::listTopics", "Goi TopicRepository::getAll va in danh sach topic cho user chon khi add flashcard."),
    ("src/TopicRepository.cpp", "TopicRepository::getAll", "SELECT danh sach topic kem subject name, map thanh vector<Topic>."),
    ("src/TopicRepository.cpp", "TopicRepository::exists", "Kiem tra TopicId co ton tai khong. Day la lop bao ve app truoc loi foreign key khi add flashcard."),
    ("src/ReviewService.cpp", "ReviewService::reviewDueFlashcards", "Lay cac flashcard den han, in cau hoi, doi user bam Enter, hien dap an, hoi dung/sai, roi luu review."),
    ("src/ReviewService.cpp", "ReviewService::showWeakTopics", "Lay danh sach topic co ty le sai cao va in ra man hinh."),
    ("src/ReviewRepository.cpp", "ReviewRepository::getDueFlashcards", "SELECT TOP 10 flashcards co NextReviewAt <= SYSUTCDATETIME, uu tien difficulty cao va den han som."),
    ("src/ReviewRepository.cpp", "ReviewRepository::saveReview", "INSERT vao ReviewLogs, sau do UPDATE Flashcards de tinh ReviewIntervalDays va NextReviewAt theo dung/sai."),
    ("src/ReviewRepository.cpp", "ReviewRepository::getWeakTopics", "Aggregate ReviewLogs theo subject/topic de tinh reviewCount, wrongCount va wrongRate."),
    ("src/StatisticsService.cpp", "StatisticsService::showStudyStatistics", "Goi Repository lay thong ke va in ra terminal."),
    ("src/StatisticsRepository.cpp", "readIntResult", "Ham phu chay cau SELECT tra ve mot so duy nhat, neu khong co ket qua thi tra 0."),
    ("src/StatisticsRepository.cpp", "StatisticsRepository::getStudyStatistics", "Gom cac query COUNT va accuracy de tao object StudyStatistics."),
    ("web/backend/app/db.py", "build_connection_string", "Tao connection string Python/pyodbc tu bien moi truong .env. Day la ban Python cua connection string ben C++."),
    ("web/backend/app/db.py", "get_connection", "Context manager mo ket noi pyodbc va dam bao dong connection sau khi dung xong."),
    ("web/backend/app/db.py", "fetch_all", "Chay SELECT va tra ve list dictionary, moi row co key la ten cot."),
    ("web/backend/app/db.py", "fetch_one", "Dung fetch_all roi lay dong dau tien, phu hop voi query chi can mot object hoac mot count."),
    ("web/backend/app/db.py", "execute", "Chay INSERT/UPDATE/DELETE va commit transaction."),
    ("web/backend/app/repositories.py", "list_flashcards", "API repository lay danh sach flashcard cho web."),
    ("web/backend/app/repositories.py", "search_flashcards", "API repository tim flashcard theo keyword bang parameter query an toan hon noi chuoi SQL."),
    ("web/backend/app/repositories.py", "flashcard_exists", "Kiem tra flashcard ton tai truoc update/delete/review tren web."),
    ("web/backend/app/repositories.py", "create_flashcard", "Them flashcard moi tu request cua frontend."),
    ("web/backend/app/repositories.py", "update_flashcard", "Cap nhat flashcard tu request PUT."),
    ("web/backend/app/repositories.py", "delete_flashcard", "Xoa ReviewLogs roi xoa Flashcards cho web."),
    ("web/backend/app/repositories.py", "list_topics", "Lay topics de frontend render dropdown chon topic."),
    ("web/backend/app/repositories.py", "topic_exists", "Validate TopicId truoc khi API tao flashcard."),
    ("web/backend/app/repositories.py", "due_reviews", "Lay danh sach card den han cho man hinh review web."),
    ("web/backend/app/repositories.py", "save_review", "Luu ket qua review va tinh lich on tiep theo tren web."),
    ("web/backend/app/repositories.py", "weak_topics", "Lay topic yeu cho dashboard web."),
    ("web/backend/app/repositories.py", "study_statistics", "Lay thong ke tong hop cho dashboard web."),
    ("web/backend/app/main.py", "index", "Tra ve file index.html khi user mo http://127.0.0.1:8000."),
    ("web/backend/app/main.py", "health", "Endpoint kiem tra server FastAPI co dang chay khong."),
    ("web/backend/app/main.py", "list_flashcards", "Route GET /api/flashcards, goi repository va tra JSON."),
    ("web/backend/app/main.py", "search_flashcards", "Route GET /api/flashcards/search, nhan keyword query string."),
    ("web/backend/app/main.py", "create_flashcard", "Route POST /api/flashcards, validate topic roi tao flashcard."),
    ("web/backend/app/main.py", "update_flashcard", "Route PUT /api/flashcards/{id}, validate flashcard roi update."),
    ("web/backend/app/main.py", "delete_flashcard", "Route DELETE /api/flashcards/{id}, validate flashcard roi xoa."),
    ("web/backend/app/main.py", "list_topics", "Route GET /api/topics cho dropdown frontend."),
    ("web/backend/app/main.py", "due_reviews", "Route GET /api/reviews/due cho man hinh review."),
    ("web/backend/app/main.py", "save_review", "Route POST /api/reviews de luu dung/sai."),
    ("web/backend/app/main.py", "weak_topics", "Route GET /api/reviews/weak-topics cho dashboard."),
    ("web/backend/app/main.py", "study_statistics", "Route GET /api/statistics/study cho dashboard."),
    ("web/frontend/app.js", "requestJson", "Ham dung chung de fetch API, check loi HTTP, parse JSON."),
    ("web/frontend/app.js", "renderStats", "Nhan statistics JSON va render cac o thong ke tren dashboard."),
    ("web/frontend/app.js", "renderFlashcards", "Nhan list flashcards va render thanh cac card HTML."),
    ("web/frontend/app.js", "renderTopics", "Nhan list topics va render option trong select."),
    ("web/frontend/app.js", "renderDueReviews", "Nhan list due reviews va render nut Correct/Wrong."),
    ("web/frontend/app.js", "renderWeakTopics", "Nhan weak topics JSON va render danh sach topic yeu."),
    ("web/frontend/app.js", "escapeForJs", "Escape dau \\ va ' truoc khi dua text vao onclick inline."),
    ("web/frontend/app.js", "loadDashboard", "Goi nhieu API song song bang Promise.all va render toan bo dashboard."),
    ("web/frontend/app.js", "searchFlashcards", "Doc keyword tu input va goi API search hoac list all neu keyword rong."),
    ("web/frontend/app.js", "addFlashcard", "Xu ly submit form, POST flashcard moi, reset form va reload dashboard."),
    ("web/frontend/app.js", "editFlashcard", "Hoi user bang prompt, validate difficulty, PUT len API, reload dashboard."),
    ("web/frontend/app.js", "deleteFlashcard", "Hoi confirm, DELETE flashcard, reload dashboard."),
    ("web/frontend/app.js", "saveReview", "POST ket qua dung/sai len API, reload dashboard."),
]


def set_paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color="DADCE0", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn("w:" + edge))
        if element is None:
            element = OxmlElement("w:" + edge)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def clean_text(text):
    allowed = []
    for ch in text:
        if ch in ("\n", "\t", "\r") or ord(ch) >= 32:
            allowed.append(ch)
    return "".join(allowed)


def add_code_block(doc, code, max_lines_per_block=90):
    lines = clean_text(code).rstrip().splitlines()
    if not lines:
        lines = [""]
    for start in range(0, len(lines), max_lines_per_block):
        chunk = lines[start:start + max_lines_per_block]
        for line in chunk:
            p = doc.add_paragraph(style="CodeLine")
            set_paragraph_shading(p, "F6F8FA")
            run = p.add_run(line if line else " ")
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        doc.add_paragraph()


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_note(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    set_table_borders(table, "BFC7D5")
    cell = table.cell(0, 0)
    set_cell_shading(cell, "EEF3FA")
    p = cell.paragraphs[0]
    p.style = doc.styles["Normal"]
    r = p.add_run(title + ": ")
    r.bold = True
    p.add_run(text)


def add_simple_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = True
    set_table_borders(table)
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_shading(cell, "E8EEF5")
        p = cell.paragraphs[0]
        p.style = doc.styles["TableText"]
        r = p.add_run(header)
        r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            p.style = doc.styles["TableText"]
            p.add_run(value)
    doc.add_paragraph()


def setup_doc():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(11, 37, 69)

    for style_name, size, color in [
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "2E74B5"),
        ("Heading 3", 11.5, "1F4D78"),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)

    code = styles.add_style("CodeLine", 1)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    code.font.size = Pt(7.5)
    code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.left_indent = Inches(0.12)
    code.paragraph_format.line_spacing = 1.0

    table_text = styles.add_style("TableText", 1)
    table_text.font.name = "Calibri"
    table_text.font.size = Pt(8.5)
    table_text.paragraph_format.space_after = Pt(0)
    table_text.paragraph_format.line_spacing = 1.05

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
    footer.add_run("VS Code codebase guide - Smart Study Planner")
    return doc


def read_file(relative_path):
    path = ROOT / relative_path
    return path.read_text(encoding="utf-8", errors="replace")


def functions_for_file(relative_path):
    return [(name, meaning) for file_path, name, meaning in FUNCTIONS if file_path == relative_path]


def build_doc():
    doc = setup_doc()

    doc.add_paragraph("VS Code Codebase Guide", style="Title")
    doc.add_paragraph("Smart Study Planner - giải thích code cho người mới bắt đầu")
    doc.add_paragraph(
        "Tài liệu này chỉ tập trung vào phần code đang nằm trong VS Code. "
        "Mục tiêu là khi mở project lên, người mới có thể hiểu từng thư mục, từng file, từng function, "
        "vì sao code được tách lớp, và ứng dụng kết nối SQL Server như thế nào."
    )

    doc.add_heading("1. VS Code trong dự án này là gì?", level=1)
    doc.add_paragraph(
        "VS Code không tự chạy database và cũng không tự hiểu SQL Server. Trong dự án này, VS Code là nơi ta mở source code, "
        "sửa file, mở terminal, chạy CMake, chạy file .exe, chạy FastAPI backend và dùng Git. "
        "Phần thật sự kết nối SQL Server nằm trong code C++/Python, thông qua ODBC driver."
    )
    add_bullets(
        doc,
        [
            "Explorer bên trái: xem folder include, src, sql, web, docs.",
            "Editor ở giữa: sửa file .cpp, .h, .sql, .py, .js, .html.",
            "Terminal bên dưới: chạy cmake, chạy app, chạy git, chạy uvicorn.",
            "Source Control: xem file nào đã sửa, commit và push.",
            "Extensions C++/CMake: hỗ trợ gợi ý code, IntelliSense và build.",
        ],
    )

    doc.add_heading("2. Cấu trúc project nhìn từ VS Code", level=1)
    add_code_block(
        doc,
        """project/
  include/                  Header files: khai báo class, struct, function
    models/                 Các struct dữ liệu như Flashcard, Topic
  src/                      Source files: code xử lý thật sự
  sql/                      Script tạo database và dữ liệu mẫu
  web/
    backend/                FastAPI backend chuẩn bị cho web
    frontend/               HTML/CSS/JS prototype
  docs/                     Tài liệu học lại, API contract, handoff
  CMakeLists.txt            File cấu hình build C++
  README.md                 Tóm tắt dự án""",
    )
    doc.add_paragraph(
        "Cách đọc quan trọng: include thường cho biết app có những lớp/hàm nào; src cho biết các hàm đó được viết như thế nào. "
        "sql cho biết database có bảng gì. web là phần chuẩn bị để sau này thay CLI bằng giao diện web."
    )

    doc.add_heading("3. Kết nối VS Code với SQL Server trong dự án này", level=1)
    doc.add_paragraph(
        "Nói chính xác hơn: không phải VS Code kết nối SQL Server, mà chương trình được chạy từ terminal VS Code kết nối SQL Server."
    )
    add_numbered(
        doc,
        [
            "SQL Server Express chạy instance HOANE\\SQLEXPRESS.",
            "Database SmartStudyPlanner đã được tạo bằng sql/schema.sql.",
            "Máy có ODBC Driver 18 for SQL Server.",
            "C++ build thành smart_study_planner.exe bằng CMake/MSVC.",
            "Khi chạy .exe trong terminal, main.cpp tạo connection string.",
            "DbConnection::connect dùng SQLDriverConnectW của ODBC để mở kết nối.",
            "Repository gửi SELECT/INSERT/UPDATE/DELETE qua DbConnection.",
            "SQL Server trả dữ liệu về, app in ra terminal hoặc xử lý tiếp.",
        ],
    )
    add_code_block(
        doc,
        """std::wstring connectionString =
    L"Driver={" + driver + L"};"
    L"Server=" + server + L";"
    L"Database=" + database + L";"
    L"Encrypt=no;"
    L"TrustServerCertificate=yes;";

if (user.empty()) {
    connectionString += L"Trusted_Connection=yes;";
} else {
    connectionString += L"UID=" + user + L";PWD=" + password + L";";
}""",
    )
    add_note(
        doc,
        "Ý nghĩa connection string",
        "Driver nói app dùng ODBC Driver 18. Server trỏ tới HOANE\\SQLEXPRESS. Database trỏ tới SmartStudyPlanner. "
        "Trusted_Connection=yes nghĩa là đăng nhập bằng tài khoản Windows hiện tại. Encrypt=no và TrustServerCertificate=yes giúp tránh lỗi chứng chỉ trong môi trường local học tập.",
    )

    doc.add_heading("4. Vì sao phải tách Model / Service / Repository?", level=1)
    doc.add_paragraph(
        "Nếu toàn bộ code viết trong main.cpp, ban đầu nhìn có vẻ nhanh, nhưng về sau rất khó sửa. "
        "Một dự án thật thường tách lớp để mỗi phần chịu một trách nhiệm rõ ràng."
    )
    add_simple_table(
        doc,
        ["Lớp", "Nhiệm vụ", "Ví dụ trong project"],
        [
            ("Model", "Mô tả hình dạng dữ liệu.", "Flashcard có id, subjectName, topicName, question, answer, difficulty."),
            ("Service", "Xử lý nghiệp vụ và tương tác user.", "Add flashcard: hỏi input, validate, gọi repository."),
            ("Repository", "Chứa SQL, làm việc trực tiếp với database.", "INSERT/UPDATE/DELETE Flashcards."),
            ("DbConnection", "Lớp thấp nhất để gửi SQL qua ODBC.", "query() và execute()."),
            ("Menu/main", "Điều hướng chương trình.", "User chọn 1 thì gọi FlashcardService::listFlashcards."),
        ],
    )
    doc.add_paragraph("Những lợi ích quan trọng khi đi làm:")
    add_bullets(
        doc,
        [
            "Dễ đọc: người mới biết tìm SQL ở Repository, tìm logic user ở Service.",
            "Dễ sửa: đổi câu SQL không cần đụng Menu; đổi giao diện không cần viết lại SQL.",
            "Dễ test: có thể test Repository riêng, Service riêng.",
            "Dễ nâng cấp web: CLI và web có thể dùng cùng ý tưởng Repository/Service.",
            "Dễ làm nhóm: một người làm UI, một người làm Service, một người làm database ít đụng file nhau.",
            "Giảm bug dây chuyền: sửa một lớp ít ảnh hưởng lớp khác.",
            "Gần với kiến trúc đi làm: nhiều dự án thật dùng Controller/Service/Repository/Model hoặc biến thể tương tự.",
        ],
    )

    doc.add_heading("5. Bản đồ function theo từng file", level=1)
    doc.add_paragraph(
        "Phần này là trọng tâm. Mỗi file có code nguyên bản đang nằm trong VS Code, sau đó là bảng giải thích từng function trong file đó."
    )

    for relative_path, description in FILES_TO_INCLUDE:
        doc.add_heading(f"{relative_path}", level=2)
        doc.add_paragraph(description)
        code = read_file(relative_path)
        add_code_block(doc, code)
        file_functions = functions_for_file(relative_path)
        if file_functions:
            add_simple_table(
                doc,
                ["Function", "Function này để làm gì?"],
                [(name, meaning) for name, meaning in file_functions],
            )
        else:
            doc.add_paragraph(
                "File này không có function xử lý chính, nhưng vẫn quan trọng vì nó khai báo dữ liệu, cấu hình build, layout hoặc class/schema."
            )

    doc.add_heading("6. Luồng chạy C++ từ lúc bấm run tới lúc có dữ liệu", level=1)
    add_numbered(
        doc,
        [
            "Terminal trong VS Code chạy .\\build\\smart_study_planner.exe.",
            "Windows gọi wmain trong src/main.cpp.",
            "wmain tạo connection string và gọi DbConnection::connect.",
            "DbConnection dùng ODBC mở kết nối tới HOANE\\SQLEXPRESS.",
            "wmain gọi showMenu để in menu.",
            "User chọn một số, switch-case trong main gọi Service tương ứng.",
            "Service đọc input và validate.",
            "Service gọi Repository.",
            "Repository tạo SQL string và gọi db.query hoặc db.execute.",
            "DbConnection gửi SQL tới SQL Server.",
            "SQL Server trả kết quả, Repository map dữ liệu thành model.",
            "Service in kết quả ra terminal.",
        ],
    )

    doc.add_heading("7. Luồng chuẩn bị web từ code hiện tại", level=1)
    doc.add_paragraph(
        "Web hiện mới là phần chuẩn bị, nhưng cách tổ chức giống CLI: backend FastAPI có db.py giống DbConnection, repositories.py giống Repository, main.py giống Controller/API routes."
    )
    add_code_block(
        doc,
        """Browser
-> frontend/index.html + app.js
-> fetch('/api/...')
-> FastAPI main.py route
-> repositories.py
-> db.py / pyodbc
-> SQL Server""",
    )
    doc.add_paragraph(
        "Điểm cần hiểu: khi chuyển lên web, người dùng không gõ menu trong terminal nữa. Họ bấm nút trên trình duyệt. "
        "Nút đó gọi JavaScript, JavaScript gọi API, API gọi SQL Server."
    )

    doc.add_heading("8. Cách tự học lại project này trong VS Code", level=1)
    add_numbered(
        doc,
        [
            "Đọc README.md để biết app làm gì.",
            "Đọc sql/schema.sql để hiểu database lưu gì.",
            "Đọc include/models để biết dữ liệu đi trong app có hình dạng gì.",
            "Đọc src/main.cpp để hiểu luồng chạy tổng quát.",
            "Đọc src/Menu.cpp để hiểu menu.",
            "Đọc từng Service để hiểu nghiệp vụ user.",
            "Đọc từng Repository để hiểu SQL.",
            "Đọc DbConnection cuối cùng nếu muốn hiểu ODBC sâu hơn.",
            "Sau đó mới đọc web/backend và web/frontend.",
        ],
    )
    add_note(
        doc,
        "Cách học hiệu quả",
        "Đừng cố hiểu tất cả cùng lúc. Hãy chọn một chức năng, ví dụ Add flashcard, rồi đi theo đường main.cpp -> FlashcardService::addFlashcard -> TopicRepository::exists -> FlashcardRepository::create -> DbConnection::execute -> SQL Server.",
    )

    doc.add_heading("9. Những lỗi VS Code/codebase hay gặp và cách nhìn đúng", level=1)
    add_simple_table(
        doc,
        ["Hiện tượng", "Cách hiểu đúng", "Cách xử lý"],
        [
            ("VS Code gạch đỏ code C++", "Có thể IntelliSense chưa nhận đúng compiler/include, chưa chắc build lỗi.", "Chạy cmake --build build để xác nhận thật sự."),
            ("cmake --build build lỗi linker", "Thường do thêm .cpp mới nhưng quên thêm vào CMakeLists.txt.", "Mở CMakeLists.txt, thêm file .cpp vào add_executable."),
            ("App báo không kết nối SQL Server", "Connection string, SQL Server service, ODBC driver hoặc database chưa đúng.", "Kiểm tra SSMS connect được, database tồn tại, driver đúng."),
            ("Add flashcard với TopicId sai", "Database chặn vì foreign key.", "Validate TopicRepository::exists trước khi insert."),
            ("Delete flashcard lỗi foreign key", "ReviewLogs còn trỏ đến Flashcards.", "Xóa ReviewLogs trước, rồi xóa Flashcards."),
            ("Python command không chạy đúng", "Windows Store alias có thể chen vào PATH.", "Dùng venv: web/backend/.venv/Scripts/python.exe."),
        ],
    )

    doc.add_heading("10. Kết luận ngắn", level=1)
    doc.add_paragraph(
        "Phần VS Code của dự án này không chỉ là vài file code. Nó là một hệ thống nhỏ có cấu trúc rõ: "
        "database lưu dữ liệu, C++ CLI xử lý nghiệp vụ, Repository chứa SQL, Service điều phối logic, "
        "DbConnection nối SQL Server, và web scaffold chuẩn bị biến CLI thành ứng dụng web. "
        "Nếu nắm được luồng này, những dự án sau như quản lý thư viện, đặt sân, bán hàng, task manager đều có thể làm theo kiểu tương tự."
    )

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_doc()
    print(OUTPUT)
