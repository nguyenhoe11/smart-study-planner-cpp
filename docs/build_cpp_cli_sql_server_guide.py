from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "CPP_CLI_SQL_Server_Express_Guide.docx"
ASSET_DIR = ROOT / "docs" / "assets" / "cpp_sql_guide"


SOURCES = [
    (
        "Microsoft Learn - ODBC API reference",
        "https://learn.microsoft.com/en-us/SQL/odbc/reference/syntax/odbc-api-reference?view=sql-server-ver15",
        "ODBC API cung cấp các hàm chuẩn để kết nối data source, chạy SQL statement và lấy kết quả.",
    ),
    (
        "Microsoft Learn - ODBC Driver for SQL Server connection string keywords",
        "https://learn.microsoft.com/en-us/sql/connect/odbc/dsn-connection-string-attribute?view=sql-server-ver17",
        "Tài liệu về Driver, Server, Database, Trusted_Connection, Encrypt, TrustServerCertificate và các keyword liên quan.",
    ),
    (
        "Microsoft Learn - C/C++ ODBC sample app",
        "https://learn.microsoft.com/en-ca/sql/connect/odbc/cpp-code-example-app-connect-access-sql-db?view=azure-sqldw-latest",
        "Ví dụ chính thức về ứng dụng C/C++ dùng ODBC API để kết nối và truy cập SQL database.",
    ),
    (
        "Microsoft Learn - SQL Server editions and supported features",
        "https://learn.microsoft.com/en-us/sql/sql-server/editions-and-components-of-sql-server-2022?view=sql-server-ver16",
        "Tài liệu tổng quan về SQL Server Database Engine và các edition, trong đó có Express.",
    ),
]


def font(size=28, bold=False):
    candidates = [
        Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\segoeuib.ttf" if bold else r"C:\Windows\Fonts\segoeui.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def rounded(draw, xy, fill, outline="#334155", radius=18, width=2):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def centered(draw, box, text, fill="#0f172a", size=24, bold=False):
    f = font(size, bold)
    lines = text.split("\n")
    line_heights = []
    widths = []
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=f)
        widths.append(bbox[2] - bbox[0])
        line_heights.append(bbox[3] - bbox[1])
    total_h = sum(line_heights) + (len(lines) - 1) * 8
    y = box[1] + ((box[3] - box[1]) - total_h) / 2
    for i, line in enumerate(lines):
        x = box[0] + ((box[2] - box[0]) - widths[i]) / 2
        draw.text((x, y), line, font=f, fill=fill)
        y += line_heights[i] + 8


def arrow(draw, start, end, color="#334155", width=4):
    draw.line([start, end], fill=color, width=width)
    x1, y1 = start
    x2, y2 = end
    if x2 >= x1:
        points = [(x2, y2), (x2 - 14, y2 - 8), (x2 - 14, y2 + 8)]
    else:
        points = [(x2, y2), (x2 + 14, y2 - 8), (x2 + 14, y2 + 8)]
    draw.polygon(points, fill=color)


def save_diagram(name, title, boxes, arrows_list, size=(1400, 760)):
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", size, "#ffffff")
    d = ImageDraw.Draw(img)
    d.text((40, 30), title, font=font(34, True), fill="#0b2545")
    for box in boxes:
        rounded(d, box["xy"], box.get("fill", "#eef2ff"), box.get("outline", "#334155"))
        centered(d, box["xy"], box["text"], size=box.get("size", 23), bold=box.get("bold", True))
    for item in arrows_list:
        arrow(d, item[0], item[1], color=item[2] if len(item) > 2 else "#334155")
    path = ASSET_DIR / f"{name}.png"
    img.save(path)
    return path


def create_diagrams():
    diagrams = {}
    diagrams["big_picture"] = save_diagram(
        "big_picture",
        "Bức tranh tổng thể: C++ CLI làm việc với SQL Server Express",
        [
            {"xy": (40, 220, 250, 360), "text": "Người dùng\nnhập menu", "fill": "#e0f2fe"},
            {"xy": (320, 220, 560, 360), "text": "C++ CLI\n.exe trong terminal", "fill": "#dcfce7"},
            {"xy": (630, 220, 850, 360), "text": "DbConnection\nODBC API", "fill": "#fef9c3"},
            {"xy": (920, 220, 1140, 360), "text": "ODBC Driver 18\nfor SQL Server", "fill": "#ffedd5"},
            {"xy": (1210, 220, 1370, 360), "text": "SQL Server\nExpress", "fill": "#fee2e2"},
            {"xy": (1040, 500, 1350, 670), "text": "Database\nSmartStudyPlanner\nSubjects / Topics /\nFlashcards / ReviewLogs", "fill": "#f1f5f9", "size": 21},
        ],
        [
            ((250, 290), (320, 290)),
            ((560, 290), (630, 290)),
            ((850, 290), (920, 290)),
            ((1140, 290), (1210, 290)),
            ((1290, 360), (1190, 500)),
        ],
    )
    diagrams["odbc_layers"] = save_diagram(
        "odbc_layers",
        "ODBC giống một bộ chuyển ngữ giữa C++ và SQL Server",
        [
            {"xy": (90, 120, 1310, 220), "text": "C++ app gọi hàm chuẩn: SQLDriverConnectW, SQLExecDirectW, SQLFetch, SQLGetData", "fill": "#dbeafe", "size": 24},
            {"xy": (180, 280, 1220, 370), "text": "ODBC Driver Manager của Windows chọn đúng driver", "fill": "#dcfce7"},
            {"xy": (260, 430, 1140, 520), "text": "ODBC Driver 18 for SQL Server hiểu giao thức SQL Server", "fill": "#fef9c3"},
            {"xy": (360, 580, 1040, 680), "text": "SQL Server Express xử lý SQL và trả kết quả", "fill": "#fee2e2"},
        ],
        [
            ((700, 220), (700, 280)),
            ((700, 370), (700, 430)),
            ((700, 520), (700, 580)),
        ],
    )
    diagrams["query_flow"] = save_diagram(
        "query_flow",
        "Một lần user chọn 'List flashcards' chạy qua những bước nào?",
        [
            {"xy": (40, 170, 250, 300), "text": "Menu\nchoice = 1", "fill": "#e0f2fe"},
            {"xy": (310, 170, 540, 300), "text": "FlashcardService\nlistFlashcards", "fill": "#dcfce7"},
            {"xy": (600, 170, 850, 300), "text": "FlashcardRepository\ngetAll", "fill": "#fef9c3"},
            {"xy": (910, 170, 1130, 300), "text": "DbConnection\nquery", "fill": "#ffedd5"},
            {"xy": (1190, 170, 1370, 300), "text": "SQL Server\nSELECT", "fill": "#fee2e2"},
            {"xy": (600, 470, 850, 620), "text": "map rows\n-> Flashcard objects", "fill": "#f1f5f9", "size": 22},
            {"xy": (310, 470, 540, 620), "text": "Service in kết quả\nra terminal", "fill": "#dcfce7"},
        ],
        [
            ((250, 235), (310, 235)),
            ((540, 235), (600, 235)),
            ((850, 235), (910, 235)),
            ((1130, 235), (1190, 235)),
            ((1280, 300), (760, 470)),
            ((600, 545), (540, 545)),
        ],
    )
    diagrams["tables"] = save_diagram(
        "tables",
        "Quan hệ dữ liệu trong database SmartStudyPlanner",
        [
            {"xy": (80, 260, 330, 420), "text": "Subjects\nSubjectId\nName", "fill": "#dbeafe"},
            {"xy": (430, 260, 680, 420), "text": "Topics\nTopicId\nSubjectId\nName", "fill": "#dcfce7"},
            {"xy": (780, 260, 1030, 420), "text": "Flashcards\nFlashcardId\nTopicId\nQuestion\nAnswer", "fill": "#fef9c3"},
            {"xy": (1130, 260, 1370, 420), "text": "ReviewLogs\nReviewLogId\nFlashcardId\nWasCorrect", "fill": "#fee2e2", "size": 21},
        ],
        [
            ((330, 340), (430, 340)),
            ((680, 340), (780, 340)),
            ((1030, 340), (1130, 340)),
        ],
    )
    diagrams["connection_string"] = save_diagram(
        "connection_string",
        "Connection string giống một tờ địa chỉ để app tìm đúng database",
        [
            {"xy": (90, 170, 360, 310), "text": "Driver\nODBC Driver 18", "fill": "#dbeafe"},
            {"xy": (410, 170, 680, 310), "text": "Server\nHOANE\\SQLEXPRESS", "fill": "#dcfce7"},
            {"xy": (730, 170, 1000, 310), "text": "Database\nSmartStudyPlanner", "fill": "#fef9c3"},
            {"xy": (1050, 170, 1320, 310), "text": "Auth\nTrusted_Connection", "fill": "#ffedd5"},
            {"xy": (250, 470, 590, 620), "text": "Encrypt / Certificate\nlocal dev setting", "fill": "#fee2e2"},
            {"xy": (760, 470, 1150, 620), "text": "Kết quả\nmở được đường tới SQL Server", "fill": "#f1f5f9"},
        ],
        [
            ((225, 310), (420, 470)),
            ((545, 310), (500, 470)),
            ((865, 310), (910, 470)),
            ((1185, 310), (1010, 470)),
            ((590, 545), (760, 545)),
        ],
    )
    return diagrams


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color="DADCE0", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn("w:" + edge))
        if element is None:
            element = OxmlElement("w:" + edge)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def setup_doc():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(11, 37, 69)

    for style_name, size, color in [
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "2E74B5"),
        ("Heading 3", 11.5, "1F4D78"),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)

    code = styles.add_style("MiniCode", 1)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    code.font.size = Pt(8.5)
    code.paragraph_format.space_after = Pt(2)

    table_text = styles.add_style("TableText", 1)
    table_text.font.name = "Calibri"
    table_text.font.size = Pt(8.5)
    table_text.paragraph_format.space_after = Pt(0)
    table_text.paragraph_format.line_spacing = 1.05

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
    footer.add_run("C++ CLI + SQL Server Express guide")
    return doc


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_note(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, "BFC7D5")
    cell = table.cell(0, 0)
    set_cell_shading(cell, "EEF3FA")
    p = cell.paragraphs[0]
    p.style = doc.styles["Normal"]
    r = p.add_run(title + ": ")
    r.bold = True
    p.add_run(text)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_shading(cell, "E8EEF5")
        p = cell.paragraphs[0]
        p.style = doc.styles["TableText"]
        r = p.add_run(header)
        r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            p.style = doc.styles["TableText"]
            p.add_run(value)
    doc.add_paragraph()


def add_mini_code(doc, text):
    p = doc.add_paragraph(style="MiniCode")
    r = p.add_run(text)
    r.font.name = "Consolas"


def add_image(doc, path, caption):
    doc.add_picture(str(path), width=Inches(6.7))
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    r = p.add_run(caption)
    r.italic = True
    r.font.size = Pt(8.5)


def build_doc():
    diagrams = create_diagrams()
    doc = setup_doc()

    doc.add_paragraph("C++ CLI is working with SQL Server Express database", style="Title")
    doc.add_paragraph("Giải thích cực kỳ chi tiết từ cơ bản đến nâng cao, theo đúng dự án Smart Study Planner")
    doc.add_paragraph(
        "Câu này nghe ngắn, nhưng bên trong có nhiều lớp kiến thức: CLI là gì, C++ chạy như một process ra sao, "
        "SQL Server Express là gì, database khác server thế nào, ODBC driver đóng vai trò gì, connection string là gì, "
        "và tại sao khi terminal in được dữ liệu thì ta nói app đã làm việc được với database."
    )

    doc.add_heading("1. Dịch câu này sang ngôn ngữ dễ hiểu", level=1)
    doc.add_paragraph(
        "Trong dự án này, câu 'C++ CLI is working with SQL Server Express database' nghĩa là: "
        "file chương trình C++ chạy trong terminal đã mở được kết nối tới SQL Server Express, chọn đúng database SmartStudyPlanner, "
        "gửi được câu lệnh SQL, nhận được dữ liệu trả về, rồi in hoặc xử lý dữ liệu đó trong menu terminal."
    )
    add_table(
        doc,
        ["Cụm từ", "Nghĩa đơn giản", "Trong dự án này"],
        [
            ("C++", "Ngôn ngữ viết chương trình chính.", "File .cpp build thành smart_study_planner.exe."),
            ("CLI", "Command Line Interface: app chạy bằng chữ trong terminal.", "Menu 1-9, user nhập số để chọn chức năng."),
            ("working with", "Không chỉ mở app, mà app thật sự gửi/nhận dữ liệu.", "List, search, add, edit, delete, review đều đụng SQL Server."),
            ("SQL Server Express", "Bản SQL Server miễn phí/nhẹ cho học tập và app nhỏ.", "Instance HOANE\\SQLEXPRESS."),
            ("database", "Kho dữ liệu cụ thể bên trong SQL Server.", "SmartStudyPlanner có Subjects, Topics, Flashcards, ReviewLogs."),
        ],
    )
    add_note(
        doc,
        "Ví dụ đời thường",
        "Hãy tưởng tượng SQL Server Express là một tòa nhà thư viện, database SmartStudyPlanner là một phòng trong thư viện, "
        "các bảng là từng kệ sách, còn C++ CLI là người thủ thư đứng ở quầy. ODBC driver là thẻ ra vào và người phiên dịch giúp thủ thư nói chuyện đúng cách với hệ thống thư viện.",
    )
    add_image(doc, diagrams["big_picture"], "Hình 1 - Luồng tổng thể từ terminal đến SQL Server Express.")

    doc.add_heading("2. CLI là gì và tại sao dùng CLI trước web?", level=1)
    doc.add_paragraph(
        "CLI là kiểu ứng dụng giao tiếp với người dùng bằng dòng lệnh. Người dùng thấy chữ, nhập số hoặc text, app xử lý rồi in kết quả. "
        "CLI không có nút bấm, không có form đẹp, không có CSS, nhưng nó rất tốt để học bản chất vì mọi thứ hiện ra trực tiếp: input, output, lỗi và dữ liệu."
    )
    add_bullets(
        doc,
        [
            "CLI giúp tập trung vào logic trước khi bị phân tâm bởi giao diện.",
            "CLI dễ test: chạy app, chọn menu, nhìn output.",
            "CLI hợp với C++ mới học vì thấy rõ luồng chương trình: main -> menu -> function.",
            "CLI là nền tốt để lên web: khi logic/database đã ổn, web chỉ là lớp giao diện mới.",
        ],
    )
    doc.add_paragraph(
        "Trong dự án Smart Study Planner, CLI có menu: list flashcards, search, add, edit, delete, review, weak topics, statistics. "
        "Mỗi lựa chọn là một hành động thật với database."
    )

    doc.add_heading("3. SQL Server Express là gì?", level=1)
    doc.add_paragraph(
        "SQL Server là hệ quản trị cơ sở dữ liệu quan hệ của Microsoft. Nó là một dịch vụ chạy nền trên Windows, chịu trách nhiệm lưu bảng, "
        "kiểm tra khóa chính/khóa ngoại, xử lý SELECT/INSERT/UPDATE/DELETE, tối ưu query và trả kết quả cho app. "
        "Express là edition nhẹ/miễn phí, phù hợp cho học tập, desktop app, prototype và ứng dụng nhỏ."
    )
    doc.add_paragraph(
        "Điểm rất dễ nhầm: SSMS không phải database. SSMS chỉ là phần mềm quản lý. SQL Server Express mới là server/database engine. "
        "App C++ không kết nối vào SSMS; app C++ kết nối vào SQL Server Express."
    )
    add_table(
        doc,
        ["Thứ", "Vai trò", "Có bắt buộc khi app chạy không?"],
        [
            ("SQL Server Express service", "Lưu và xử lý dữ liệu.", "Có. Service phải Running."),
            ("Database SmartStudyPlanner", "Nơi chứa bảng của dự án.", "Có. App query vào database này."),
            ("SSMS", "Giao diện quản lý database.", "Không. Dùng để tạo bảng, xem dữ liệu, debug."),
            ("VS Code", "Nơi sửa code và chạy terminal.", "Không bắt buộc ở runtime, nhưng dùng khi phát triển."),
            ("ODBC Driver", "Cầu nối cho C++ nói chuyện với SQL Server.", "Có. App cần driver để kết nối."),
        ],
    )

    doc.add_heading("4. Server, instance, database, table khác nhau thế nào?", level=1)
    doc.add_paragraph(
        "Khi mới học, bốn từ này rất dễ bị trộn. Trong dự án này cần phân biệt rõ:"
    )
    add_table(
        doc,
        ["Khái niệm", "Nghĩa", "Ví dụ trong máy của bạn"],
        [
            ("Server/máy chủ", "Máy hoặc service đang chạy SQL Server.", "Máy Windows của bạn."),
            ("Instance", "Một bản SQL Server cụ thể đang chạy trên máy.", "HOANE\\SQLEXPRESS."),
            ("Database", "Một kho dữ liệu riêng nằm trong instance.", "SmartStudyPlanner."),
            ("Table", "Bảng dữ liệu trong database.", "Flashcards, Topics, ReviewLogs."),
            ("Row", "Một dòng dữ liệu trong bảng.", "Một flashcard cụ thể."),
            ("Column", "Một thuộc tính của bảng.", "Question, Answer, Difficulty."),
        ],
    )
    add_image(doc, diagrams["tables"], "Hình 2 - Quan hệ bảng trong SmartStudyPlanner.")
    doc.add_paragraph(
        "Nếu app kết nối đúng server nhưng sai database, query dbo.Flashcards có thể báo không tìm thấy bảng. "
        "Nếu đúng database nhưng SQL Server service chưa chạy, app không mở được kết nối. Nếu đúng hết nhưng driver thiếu, ODBC không biết cách nói chuyện với SQL Server."
    )

    doc.add_heading("5. 'Working with database' được chứng minh bằng gì?", level=1)
    doc.add_paragraph(
        "Một app được xem là làm việc được với database khi nó không chỉ chạy lên, mà còn hoàn thành được vòng đời dữ liệu. "
        "Trong dự án này, có thể kiểm tra theo 5 mức:"
    )
    add_numbered(
        doc,
        [
            "Connect được: app in 'Connected to SQL Server: HOANE\\SQLEXPRESS'.",
            "Read được: app SELECT COUNT(*) hoặc list flashcards.",
            "Create được: app add flashcard mới vào bảng Flashcards.",
            "Update được: app edit flashcard hoặc update NextReviewAt sau review.",
            "Delete được: app xóa ReviewLogs rồi xóa Flashcards.",
        ],
    )
    add_note(
        doc,
        "Tư duy debug",
        "Nếu chỉ connect được mà list không ra, vấn đề có thể nằm ở database/table/query. Nếu list được mà add lỗi, vấn đề có thể nằm ở validate, foreign key hoặc INSERT. "
        "Nếu add được mà review lỗi, vấn đề có thể nằm ở update logic hoặc ReviewLogs.",
    )

    doc.add_heading("6. ODBC là gì? Vì sao C++ cần ODBC?", level=1)
    doc.add_paragraph(
        "C++ không tự nhiên biết cách nói chuyện với SQL Server. Muốn gửi query tới SQL Server, C++ cần một API/driver. "
        "Trong dự án này ta dùng ODBC. ODBC là một chuẩn kết nối database: app gọi các hàm ODBC chuẩn, driver cụ thể sẽ chuyển lời gọi đó thành giao thức mà database hiểu."
    )
    add_image(doc, diagrams["odbc_layers"], "Hình 3 - Các lớp ODBC giữa C++ và SQL Server.")
    doc.add_paragraph(
        "Theo tài liệu Microsoft, ODBC API cung cấp các hàm chuẩn để kết nối nguồn dữ liệu, chạy SQL statement và lấy kết quả. "
        "Trong project này các ý chính tương ứng là: mở connection, execute SQL, fetch row, get column data."
    )
    add_table(
        doc,
        ["Lớp", "Nhiệm vụ", "Nếu thiếu thì sao?"],
        [
            ("C++ app", "Gọi function ODBC và xử lý logic.", "Không có chương trình để chạy."),
            ("ODBC API", "Bộ hàm chuẩn như SQLDriverConnectW, SQLExecDirectW.", "C++ không có ngôn ngữ chung để gọi database."),
            ("ODBC Driver Manager", "Chọn driver phù hợp dựa trên connection string.", "Không tìm được driver/data source."),
            ("ODBC Driver 18 for SQL Server", "Hiểu cách kết nối SQL Server.", "Lỗi driver not found hoặc connection fail."),
            ("SQL Server Express", "Xử lý SQL và trả kết quả.", "Không có nơi lưu/truy vấn dữ liệu."),
        ],
    )

    doc.add_heading("7. Connection string là gì?", level=1)
    doc.add_paragraph(
        "Connection string là chuỗi cấu hình nói cho ODBC biết: dùng driver nào, đi tới server nào, chọn database nào, xác thực kiểu gì, "
        "và xử lý encryption/certificate thế nào. Nó giống một tờ địa chỉ giao hàng: thiếu số nhà hoặc sai tên đường thì không tới đúng nơi."
    )
    add_image(doc, diagrams["connection_string"], "Hình 4 - Các thành phần của connection string.")
    add_table(
        doc,
        ["Thành phần", "Ý nghĩa", "Trong dự án"],
        [
            ("Driver", "Tên ODBC driver cần dùng.", "ODBC Driver 18 for SQL Server."),
            ("Server", "Tên SQL Server instance.", "HOANE\\SQLEXPRESS."),
            ("Database", "Database app muốn dùng.", "SmartStudyPlanner."),
            ("Trusted_Connection", "Dùng Windows Authentication.", "yes nếu không nhập user/password."),
            ("UID/PWD", "Dùng SQL Login nếu có.", "Chỉ thêm khi user truyền --user và --password."),
            ("Encrypt", "Thiết lập encryption cho connection.", "no trong local dev để giảm lỗi certificate."),
            ("TrustServerCertificate", "Tin certificate server mà không validate đầy đủ.", "yes để tránh lỗi cert local/self-signed."),
        ],
    )
    doc.add_paragraph(
        "Microsoft có tài liệu riêng về connection string keyword của ODBC Driver for SQL Server, bao gồm Encrypt và TrustServerCertificate. "
        "Trong môi trường học local, cấu hình hiện tại giúp dễ chạy. Trong môi trường công ty/production, phần encryption/certificate cần được cấu hình nghiêm túc hơn."
    )

    doc.add_heading("8. Windows Authentication hoạt động thế nào?", level=1)
    doc.add_paragraph(
        "Trong project này, khi connection string có Trusted_Connection=yes, app không gửi username/password SQL riêng. "
        "SQL Server kiểm tra tài khoản Windows đang chạy app. Nghĩa là nếu bạn đang đăng nhập Windows bằng HOANE\\Admin và tài khoản đó có quyền trong SQL Server, app sẽ vào được."
    )
    add_bullets(
        doc,
        [
            "Ưu điểm: không phải hard-code password trong code.",
            "Dễ dùng cho project local cá nhân.",
            "Phù hợp khi SQL Server và app chạy cùng máy.",
            "Nếu đổi sang server thật hoặc user khác, cần cấp quyền SQL Server cho user đó.",
        ],
    )
    add_note(
        doc,
        "So sánh nhanh",
        "Windows Authentication giống dùng thẻ sinh viên để vào thư viện. SQL Login giống dùng tài khoản/password riêng do thư viện cấp.",
    )

    doc.add_heading("9. DbConnection là lớp gì trong kiến trúc?", level=1)
    doc.add_paragraph(
        "DbConnection là lớp bọc ODBC API. Thay vì để mọi repository tự gọi SQLAllocHandle, SQLExecDirectW, SQLFetch, SQLGetData, "
        "ta gom những chi tiết khó đó vào một class. Các phần khác chỉ cần gọi query hoặc execute."
    )
    add_table(
        doc,
        ["Function trong DbConnection", "Nhiệm vụ ở mức ý tưởng", "Loại SQL thường đi qua"],
        [
            ("connect", "Mở đường kết nối tới SQL Server.", "Connection string."),
            ("disconnect", "Đóng kết nối.", "Không chạy SQL."),
            ("query", "Chạy SQL có kết quả trả về.", "SELECT."),
            ("execute", "Chạy SQL không cần bảng kết quả.", "INSERT, UPDATE, DELETE."),
            ("throwIfFailed", "Biến lỗi ODBC thành exception dễ đọc hơn.", "Mọi thao tác ODBC."),
        ],
    )
    doc.add_paragraph(
        "Tư duy quan trọng: DbConnection không biết flashcard là gì. Nó chỉ biết gửi SQL và nhận dữ liệu. "
        "Điều này giúp DbConnection tái sử dụng được cho nhiều bảng khác nhau."
    )

    doc.add_heading("10. query và execute khác nhau thế nào?", level=1)
    doc.add_paragraph(
        "Trong app database, có hai kiểu thao tác chính: thao tác cần lấy dữ liệu trả về và thao tác chỉ làm thay đổi dữ liệu."
    )
    add_table(
        doc,
        ["Kiểu", "Dùng khi", "Ví dụ", "Kết quả mong đợi"],
        [
            ("query", "Muốn lấy dữ liệu.", "SELECT * FROM Flashcards", "Một bảng kết quả gồm dòng/cột."),
            ("execute", "Muốn thay đổi dữ liệu.", "INSERT/UPDATE/DELETE", "Không cần bảng kết quả, chỉ cần thành công hoặc lỗi."),
        ],
    )
    doc.add_paragraph(
        "Trong dự án, list/search/statistics dùng query. Add/edit/delete/saveReview dùng execute. "
        "Nếu nhầm query/execute, app vẫn có thể compile nhưng luồng xử lý dữ liệu sẽ sai."
    )

    doc.add_heading("11. Một lần List flashcards chạy qua những bước nào?", level=1)
    add_image(doc, diagrams["query_flow"], "Hình 5 - Luồng chạy một chức năng đọc dữ liệu.")
    doc.add_paragraph(
        "Khi user chọn 1 trong menu, main.cpp không tự viết SQL. main.cpp gọi FlashcardService::listFlashcards. "
        "Service gọi FlashcardRepository::getAll. Repository viết SQL JOIN các bảng. DbConnection::query gửi SQL qua ODBC. "
        "SQL Server trả result set. Repository map từng dòng thành object Flashcard. Service in danh sách ra terminal."
    )
    add_note(
        doc,
        "Vì sao phải đi vòng qua nhiều lớp?",
        "Nhìn có vẻ dài, nhưng nó giúp mỗi lớp chỉ làm một việc. Khi chuyển sang web, ta không cần viết lại SQL list flashcards từ đầu; chỉ cần API gọi repository tương tự.",
    )

    doc.add_heading("12. Result set là gì?", level=1)
    doc.add_paragraph(
        "Result set là bảng kết quả SQL Server trả về sau một câu SELECT. Nó giống một bảng tạm chỉ tồn tại trong lúc app đang đọc dữ liệu."
    )
    add_table(
        doc,
        ["FlashcardId", "SubjectName", "TopicName", "Question", "Difficulty"],
        [
            ("1", "Database", "SQL Join", "JOIN dùng để làm gì?", "3"),
            ("2", "C++", "ODBC", "Connection string là gì?", "4"),
            ("3", "Git", "Commit", "git commit dùng để làm gì?", "2"),
        ],
    )
    doc.add_paragraph(
        "ODBC không tự biến result set thành object C++. App phải đọc từng row, từng column. "
        "Trong project, mapFlashcards làm việc này: row[0] thành id, row[1] thành subjectName, row[2] thành topicName, v.v."
    )

    doc.add_heading("13. Vì sao dùng std::wstring và hàm có chữ W?", level=1)
    doc.add_paragraph(
        "Dự án dùng std::wstring và các hàm ODBC bản wide-character như SQLDriverConnectW, SQLExecDirectW. "
        "Lý do là dữ liệu học tập có thể chứa Unicode: tiếng Việt, dấu tiếng Việt, ký tự đặc biệt. "
        "Nếu dùng string/char không cẩn thận, dữ liệu tiếng Việt có thể bị lỗi font hoặc mất dấu."
    )
    add_bullets(
        doc,
        [
            "std::wstring lưu ký tự wide trên Windows.",
            "SQLWCHAR là kiểu wide-character của ODBC.",
            "NVARCHAR trong SQL Server lưu Unicode.",
            "N'...' trong SQL giúp SQL Server hiểu chuỗi Unicode.",
        ],
    )
    doc.add_paragraph(
        "Đây là lý do trong SQL insert/search có dạng N'text'. Nếu bỏ N, một số môi trường có thể xử lý Unicode không đúng."
    )

    doc.add_heading("14. SQL Server xử lý dữ liệu ra sao?", level=1)
    doc.add_paragraph(
        "Khi nhận SQL, SQL Server không chỉ đọc text rồi trả dữ liệu. Nó có nhiều bước bên trong:"
    )
    add_numbered(
        doc,
        [
            "Parse: kiểm tra câu SQL có đúng cú pháp không.",
            "Bind/validate: kiểm tra bảng, cột, quyền truy cập có tồn tại không.",
            "Optimize: chọn cách chạy query hiệu quả, ví dụ dùng index nào.",
            "Execute: đọc/ghi dữ liệu thật.",
            "Return: trả result set hoặc thông báo số dòng bị ảnh hưởng.",
        ],
    )
    doc.add_paragraph(
        "Ví dụ app gọi search flashcards. SQL Server có thể dùng index hoặc scan bảng tùy query. "
        "Với dự án nhỏ, tốc độ chưa phải vấn đề lớn; nhưng hiểu bước optimize giúp sau này học index, execution plan và performance."
    )

    doc.add_heading("15. Khóa chính, khóa ngoại giúp gì cho app?", level=1)
    doc.add_paragraph(
        "Database không chỉ là nơi cất dữ liệu. Nó còn bảo vệ tính đúng đắn của dữ liệu bằng constraint."
    )
    add_table(
        doc,
        ["Constraint", "Ý nghĩa", "Ví dụ trong project", "Nếu không có thì sao?"],
        [
            ("PRIMARY KEY", "Mỗi dòng có ID duy nhất.", "FlashcardId.", "Không phân biệt được card nào với card nào."),
            ("FOREIGN KEY", "Dòng này phải tham chiếu dòng có thật ở bảng khác.", "Flashcards.TopicId -> Topics.TopicId.", "Có thể tạo flashcard thuộc topic không tồn tại."),
            ("UNIQUE", "Không cho trùng giá trị.", "Subjects.Name.", "Có thể có nhiều môn trùng tên vô tình."),
            ("CHECK", "Giới hạn miền giá trị.", "Difficulty BETWEEN 1 AND 5.", "Difficulty có thể thành 999 hoặc -1."),
        ],
    )
    doc.add_paragraph(
        "Khi app add flashcard với TopicId sai, SQL Server chặn vì foreign key. Đó không phải SQL Server khó tính; đó là database đang bảo vệ dữ liệu."
    )

    doc.add_heading("16. Vì sao xóa ReviewLogs trước Flashcards?", level=1)
    doc.add_paragraph(
        "ReviewLogs phụ thuộc vào Flashcards. Mỗi review log phải thuộc về một flashcard có thật. "
        "Nếu xóa flashcard trước, các review log sẽ mồ côi: chúng trỏ tới FlashcardId không còn tồn tại. SQL Server không cho điều đó xảy ra."
    )
    add_numbered(
        doc,
        [
            "User chọn delete flashcard #5.",
            "App kiểm tra flashcard #5 có tồn tại.",
            "App hỏi confirm y/n.",
            "Repository xóa ReviewLogs WHERE FlashcardId = 5.",
            "Repository xóa Flashcards WHERE FlashcardId = 5.",
        ],
    )
    add_note(
        doc,
        "Bài học đi làm",
        "Khi xóa dữ liệu có quan hệ cha-con, luôn xem bảng nào phụ thuộc bảng nào. Trong hệ thống thật, người ta còn dùng transaction để đảm bảo xóa một cụm dữ liệu thành công hoặc rollback toàn bộ nếu có lỗi.",
    )

    doc.add_heading("17. Transaction là gì và project hiện tại đang ở mức nào?", level=1)
    doc.add_paragraph(
        "Transaction là một nhóm thao tác database được xem như một đơn vị. Hoặc tất cả thành công, hoặc tất cả bị hủy. "
        "Project hiện tại dùng execute từng câu riêng lẻ, đủ cho giai đoạn học. Nhưng với delete hoặc review, về lâu dài nên dùng transaction."
    )
    add_table(
        doc,
        ["Tình huống", "Không có transaction", "Có transaction"],
        [
            ("Delete flashcard", "Xóa ReviewLogs xong, xóa Flashcards lỗi thì dữ liệu có thể ở trạng thái lưng chừng.", "Nếu lỗi, rollback lại như chưa xóa gì."),
            ("Save review", "Insert ReviewLogs xong, update Flashcards lỗi thì lịch ôn không khớp log.", "Insert và update đi chung một transaction."),
        ],
    )
    doc.add_paragraph(
        "Khi đi làm, transaction rất quan trọng trong các nghiệp vụ tiền, đơn hàng, kho hàng, lịch sử giao dịch. "
        "Với dự án này, đây là hướng nâng cấp tốt sau khi CLI/web ổn định."
    )

    doc.add_heading("18. Error handling: app biết lỗi SQL bằng cách nào?", level=1)
    doc.add_paragraph(
        "ODBC function thường trả về SQLRETURN. Nếu thành công thì SQL_SUCCESS hoặc SQL_SUCCESS_WITH_INFO. Nếu lỗi, app phải đọc diagnostic record để biết lỗi cụ thể. "
        "Trong project, throwIfFailed gom logic này lại. Khi lỗi, nó ném std::runtime_error để main catch và in ra terminal."
    )
    add_bullets(
        doc,
        [
            "Driver không tồn tại: thường lỗi ở connection string/driver name.",
            "Server không tìm thấy: SQL Server service, instance name hoặc network protocol có vấn đề.",
            "Login failed: user Windows/SQL Login không có quyền.",
            "Invalid object name: sai database hoặc bảng chưa tạo.",
            "Foreign key violation: dữ liệu liên kết không hợp lệ.",
        ],
    )

    doc.add_heading("19. Những lỗi đã gặp trong dự án này, hiểu từ gốc", level=1)
    add_table(
        doc,
        ["Lỗi/hiện tượng", "Gốc vấn đề", "Cách hiểu bản chất"],
        [
            ("Không biết để master hay database", "Database chưa tạo hoặc đang cần chạy CREATE DATABASE.", "master là database hệ thống để bắt đầu; sau khi tạo SmartStudyPlanner thì chuyển sang database dự án."),
            ("Configuration Manager không thấy item", "Đang chọn node chưa đúng hoặc view chưa refresh.", "SQL Server service/protocol nằm ở node cụ thể, không phải mọi node đều có item."),
            ("C++ build lỗi", "MSVC env hoặc CMakeLists chưa đúng.", "Build C++ cần compiler, include path, lib path và danh sách .cpp đầy đủ."),
            ("App không kết nối SQL", "Một mắt xích Driver/Server/Database/Auth/Encrypt sai.", "Connection string là bản đồ; sai một trường có thể không tới được database."),
            ("Add TopicId sai", "Foreign key bảo vệ dữ liệu.", "App cần validate trước để lỗi thân thiện hơn thay vì để SQL Server ném lỗi."),
            ("Ctrl+S thấy gạch đỏ", "IntelliSense không đồng bộ với build thật.", "Ưu tiên kiểm tra bằng cmake --build build."),
        ],
    )

    doc.add_heading("20. Mức nâng cao: tại sao Repository là chỗ đúng để chứa SQL?", level=1)
    doc.add_paragraph(
        "Repository là lớp đại diện cho nguồn dữ liệu. Trong project này nguồn dữ liệu là SQL Server. "
        "Nếu để SQL rải khắp Service và main, sau này sửa rất mệt. Khi gom SQL vào Repository, ta tạo được biên giới rõ ràng."
    )
    add_bullets(
        doc,
        [
            "Service không cần biết câu SELECT JOIN chi tiết.",
            "Menu không cần biết database có bảng gì.",
            "Khi chuyển CLI sang web, API route có thể gọi repository giống CLI.",
            "Nếu sau này đổi database hoặc thêm cache, chỉ cần sửa tầng data access nhiều hơn là sửa toàn app.",
            "Đây là tư duy gần với kiến trúc production: Controller/UI -> Service -> Repository -> Database.",
        ],
    )

    doc.add_heading("21. Mức nâng cao: bảo mật connection và SQL injection", level=1)
    doc.add_paragraph(
        "Ở bản C++ hiện tại, một số query được tạo bằng cách nối chuỗi SQL. Dự án có escapeSql để tránh lỗi dấu nháy đơn, "
        "nhưng trong hệ thống thật, cách tốt hơn là parameterized query/prepared statement. "
        "Ở bản FastAPI Python scaffold, query đã dùng dấu ? parameter với pyodbc cho nhiều chỗ."
    )
    add_table(
        doc,
        ["Cách", "Ưu điểm", "Nhược điểm"],
        [
            ("Nối chuỗi SQL", "Dễ hiểu khi mới học.", "Dễ lỗi quote và có rủi ro SQL injection nếu không xử lý kỹ."),
            ("escapeSql", "Giảm lỗi với dấu nháy đơn.", "Chưa phải giải pháp mạnh nhất."),
            ("Parameterized query", "An toàn và chuyên nghiệp hơn.", "Cần học thêm cách bind parameter trong ODBC."),
        ],
    )
    doc.add_paragraph(
        "Khi đi làm, hãy ưu tiên parameterized query. Đây là một nâng cấp rất đáng làm cho bản C++ sau này."
    )

    doc.add_heading("22. Mức nâng cao: performance và index", level=1)
    doc.add_paragraph(
        "Database nhỏ thì query nào cũng có vẻ nhanh. Nhưng khi dữ liệu tăng, SQL Server cần index để tìm nhanh. "
        "Trong schema hiện tại có index trên NextReviewAt và ReviewLogs theo FlashcardId/ReviewedAt. Điều này phục vụ hai luồng hay dùng:"
    )
    add_bullets(
        doc,
        [
            "Tìm flashcard đến hạn review: WHERE NextReviewAt <= SYSUTCDATETIME().",
            "Xem lịch sử review của một flashcard: ReviewLogs theo FlashcardId và thời gian review.",
        ],
    )
    doc.add_paragraph(
        "Index giống mục lục sách. Không có mục lục, muốn tìm một chủ đề phải lật từng trang. Có mục lục, ta nhảy nhanh tới đúng trang. "
        "Nhưng index cũng có chi phí: INSERT/UPDATE/DELETE phải cập nhật thêm mục lục."
    )

    doc.add_heading("23. Mức nâng cao: app hiện tại khác gì app công ty?", level=1)
    add_table(
        doc,
        ["Khía cạnh", "Dự án hiện tại", "Khi đi làm thường cần thêm"],
        [
            ("Connection", "Một connection đơn giản trong CLI.", "Connection pooling, config theo environment."),
            ("SQL", "Một số query nối chuỗi.", "Parameterized query, stored procedure hoặc ORM tùy stack."),
            ("Transaction", "Chưa gom transaction rõ cho các thao tác nhiều bước.", "Transaction cho nghiệp vụ quan trọng."),
            ("Logging", "In lỗi ra terminal.", "Log file, monitoring, trace id."),
            ("Testing", "Test thủ công qua menu.", "Unit test, integration test, test database."),
            ("Security", "Local Windows Auth.", "Quyền tối thiểu, secret manager, TLS/certificate đúng chuẩn."),
            ("Deployment", "Chạy local.", "Server, CI/CD, config production/staging."),
        ],
    )

    doc.add_heading("24. Cách tự giải thích lại bằng một câu chuẩn portfolio", level=1)
    doc.add_paragraph(
        "Nếu đưa vào CV/GitHub, có thể viết:"
    )
    add_mini_code(
        doc,
        "Built a C++20 command-line flashcard planner connected to SQL Server Express via ODBC, with CRUD operations, review scheduling, and study statistics."
    )
    doc.add_paragraph(
        "Dịch ý: mình xây một app C++ chạy terminal, kết nối SQL Server Express bằng ODBC, có thêm/sửa/xóa/tìm flashcard, lịch ôn tập và thống kê học tập. "
        "Câu này nghe chuyên nghiệp vì nó nêu cả ngôn ngữ, kiểu app, database, công nghệ kết nối và tính năng."
    )

    doc.add_heading("25. Checklist hiểu thật sự", level=1)
    doc.add_paragraph("Nếu trả lời được các câu này, nghĩa là bạn đã hiểu phần 'C++ CLI working with SQL Server Express database':")
    add_numbered(
        doc,
        [
            "CLI khác GUI/web ở điểm nào?",
            "SQL Server Express khác SSMS ở điểm nào?",
            "HOANE\\SQLEXPRESS là server, database hay table?",
            "SmartStudyPlanner là gì?",
            "ODBC Driver 18 nằm ở đâu trong luồng kết nối?",
            "Connection string gồm những thành phần nào?",
            "Trusted_Connection=yes nghĩa là gì?",
            "query khác execute thế nào?",
            "Repository làm gì mà Service không nên làm?",
            "Foreign key giúp phát hiện lỗi nào?",
            "Vì sao xóa ReviewLogs trước Flashcards?",
            "Nếu app báo connected nhưng list flashcards lỗi, nên kiểm tra đâu?",
            "Nếu sau này chuyển lên web, phần nào có thể tái sử dụng về mặt tư duy?",
        ],
    )

    doc.add_heading("26. Nguồn tham khảo", level=1)
    for name, url, note in SOURCES:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(name + ": ").bold = True
        p.add_run(url)
        p2 = doc.add_paragraph()
        p2.paragraph_format.left_indent = Inches(0.25)
        p2.add_run(note)

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_doc()
    print(OUTPUT)
