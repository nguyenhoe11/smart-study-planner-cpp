from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Search_Fix_Guide.docx"
ASSET_DIR = ROOT / "docs" / "assets" / "search_fix"


def font(size=24, bold=False):
    candidates = [
        Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\segoeuib.ttf" if bold else r"C:\Windows\Fonts\segoeui.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def rounded(draw, xy, fill, outline="#334155", width=2, radius=18):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def text_center(draw, xy, text, size=23, bold=True, fill="#0f172a"):
    f = font(size, bold)
    lines = text.split("\n")
    heights = []
    widths = []
    for line in lines:
        box = draw.textbbox((0, 0), line, font=f)
        widths.append(box[2] - box[0])
        heights.append(box[3] - box[1])
    total_h = sum(heights) + (len(lines) - 1) * 8
    y = xy[1] + ((xy[3] - xy[1]) - total_h) / 2
    for i, line in enumerate(lines):
        x = xy[0] + ((xy[2] - xy[0]) - widths[i]) / 2
        draw.text((x, y), line, font=f, fill=fill)
        y += heights[i] + 8


def arrow(draw, start, end, color="#475467"):
    draw.line([start, end], fill=color, width=4)
    x1, y1 = start
    x2, y2 = end
    points = [(x2, y2), (x2 - 14, y2 - 8), (x2 - 14, y2 + 8)] if x2 >= x1 else [(x2, y2), (x2 + 14, y2 - 8), (x2 + 14, y2 + 8)]
    draw.polygon(points, fill=color)


def make_flow_diagram(filename, title, boxes, arrows_list):
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", (1400, 760), "#ffffff")
    draw = ImageDraw.Draw(img)
    draw.text((42, 30), title, font=font(34, True), fill="#0b2545")
    for box in boxes:
        rounded(draw, box["xy"], box.get("fill", "#eef2ff"), box.get("outline", "#334155"))
        text_center(draw, box["xy"], box["text"], size=box.get("size", 22), bold=box.get("bold", True))
    for item in arrows_list:
        arrow(draw, item[0], item[1], color=item[2] if len(item) > 2 else "#475467")
    path = ASSET_DIR / filename
    img.save(path)
    return path


def make_mock_ui():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", (1400, 780), "#f4f7fb")
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, 270, 780), fill="#101828")
    d.text((38, 38), "SS", font=font(26, True), fill="#ffffff")
    d.text((82, 34), "Smart Study Planner", font=font(25, True), fill="#ffffff")
    d.text((82, 68), "SQL SERVER STUDY WORKSPACE", font=font(17, True), fill="#98a2b3")
    for i, item in enumerate(["Dashboard", "Flashcards", "Editor", "Review", "Insights"]):
        d.text((42, 150 + i * 56), item, font=font(20), fill="#d0d5dd")

    d.text((320, 48), "Study control center", font=font(34, True), fill="#101828")
    rounded(d, (1120, 32, 1290, 78), "#e8f7ee", "#a7e0b8")
    text_center(d, (1120, 32, 1290, 78), "API connected", size=17, fill="#15803d")
    d.text((320, 126), "LIBRARY", font=font(18, True), fill="#98a2b3")
    d.text((320, 154), "Flashcards", font=font(27, True), fill="#101828")
    rounded(d, (455, 126, 925, 178), "#ffffff", "#b8c4d6", radius=8)
    d.text((475, 141), "join", font=font(23), fill="#101828")
    rounded(d, (940, 126, 1060, 178), "#2563eb", "#2563eb", radius=8)
    text_center(d, (940, 126, 1060, 178), "Search", size=20, fill="#ffffff")
    rounded(d, (1072, 126, 1175, 178), "#ffffff", "#b8c4d6", radius=8)
    text_center(d, (1072, 126, 1175, 178), "Clear", size=20)
    d.text((320, 205), 'Showing 3 results for "join".', font=font(21), fill="#667085")

    cards = [
        ("#1 Binary search hoạt động khi nào?", "Database / SQL Joins", "Difficulty 4"),
        ("#5 3", "Database / SQL Joins", "Difficulty 3"),
        ("#6 INNER JOIN khác LEFT JOIN ở điểm nào?", "Database / SQL Joins", "Difficulty 3"),
    ]
    y = 250
    for title, meta, diff in cards:
        rounded(d, (320, y, 1180, y + 122), "#ffffff", "#d8e0ec", radius=10)
        d.text((342, y + 24), title, font=font(23, True), fill="#101828")
        d.text((342, y + 58), meta, font=font(19), fill="#667085")
        rounded(d, (940, y + 26, 1078, y + 64), "#fff4df", "#f3c97a", radius=18)
        text_center(d, (940, y + 26, 1078, y + 64), diff, size=16, fill="#b45309")
        y += 148

    path = ASSET_DIR / "search_result_mock.png"
    img.save(path)
    return path


def create_images():
    old_flow = make_flow_diagram(
        "old_search_flow.png",
        "Luồng search cũ: phụ thuộc API và reload toàn dashboard",
        [
            {"xy": (50, 210, 270, 350), "text": "User nhập keyword\nvà bấm Search", "fill": "#dbeafe"},
            {"xy": (340, 210, 570, 350), "text": "searchFlashcards()\nset currentSearch", "fill": "#dcfce7"},
            {"xy": (640, 210, 870, 350), "text": "loadDashboard()\nreload mọi thứ", "fill": "#fef9c3"},
            {"xy": (940, 210, 1220, 350), "text": "GET /api/flashcards/search\nnetwork request", "fill": "#ffedd5"},
            {"xy": (420, 500, 760, 650), "text": "state.flashcards\nbị thay bằng search results", "fill": "#fee2e2"},
            {"xy": (840, 500, 1230, 650), "text": "Không có dòng báo kết quả\nkhông có live feedback", "fill": "#f1f5f9"},
        ],
        [
            ((270, 280), (340, 280)),
            ((570, 280), (640, 280)),
            ((870, 280), (940, 280)),
            ((1080, 350), (610, 500)),
            ((610, 350), (1030, 500)),
        ],
    )
    new_flow = make_flow_diagram(
        "new_search_flow.png",
        "Luồng search mới: dữ liệu gốc ổn định, lọc ngay trên frontend",
        [
            {"xy": (50, 180, 280, 320), "text": "loadDashboard()\nGET all flashcards", "fill": "#dbeafe"},
            {"xy": (350, 180, 620, 320), "text": "state.flashcards\nluôn là danh sách gốc", "fill": "#dcfce7"},
            {"xy": (690, 180, 960, 320), "text": "User gõ keyword\ninput event", "fill": "#fef9c3"},
            {"xy": (1030, 180, 1320, 320), "text": "visibleFlashcards()\nfilter local", "fill": "#ffedd5"},
            {"xy": (500, 500, 790, 650), "text": "renderFlashcards(filtered)", "fill": "#e8f0ff"},
            {"xy": (870, 500, 1240, 650), "text": "renderSearchSummary()\nShowing N results", "fill": "#e8f7ee"},
        ],
        [
            ((280, 250), (350, 250)),
            ((620, 250), (690, 250)),
            ((960, 250), (1030, 250)),
            ((1130, 320), (650, 500)),
            ((1160, 320), (1060, 500)),
        ],
    )
    root_cause = make_flow_diagram(
        "root_cause.png",
        "Gốc lỗi: trộn 'nguồn dữ liệu thật' với 'kết quả đang hiển thị'",
        [
            {"xy": (70, 160, 430, 320), "text": "Cần có\nAll flashcards\nnguồn dữ liệu thật", "fill": "#dcfce7"},
            {"xy": (520, 160, 880, 320), "text": "Cần có\nVisible flashcards\nkết quả đang nhìn thấy", "fill": "#dbeafe"},
            {"xy": (970, 160, 1320, 320), "text": "Code cũ dùng chung\nstate.flashcards\ncho cả hai vai trò", "fill": "#fee2e2"},
            {"xy": (260, 500, 1120, 650), "text": "Fix: state.flashcards giữ data gốc, visibleFlashcards() tạo danh sách hiển thị", "fill": "#fef9c3", "size": 25},
        ],
        [
            ((430, 240), (520, 240)),
            ((880, 240), (970, 240)),
            ((1140, 320), (760, 500), "#b42318"),
        ],
    )
    return {
        "old_flow": old_flow,
        "new_flow": new_flow,
        "root_cause": root_cause,
        "mock_ui": make_mock_ui(),
    }


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color="DADCE0", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn("w:" + edge))
        if element is None:
            element = OxmlElement("w:" + edge)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def setup_doc():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(11, 37, 69)

    for style_name, size, color in [
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "2E74B5"),
        ("Heading 3", 11.5, "1F4D78"),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)

    code = styles.add_style("CodeBlock", 1)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    code.font.size = Pt(7.5)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1.0

    table_text = styles.add_style("TableText", 1)
    table_text.font.name = "Calibri"
    table_text.font.size = Pt(8.5)
    table_text.paragraph_format.space_after = Pt(0)
    table_text.paragraph_format.line_spacing = 1.05

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
    footer.add_run("Smart Study Planner - Search fix guide")
    return doc


def add_code(doc, code, label=None):
    if label:
        p = doc.add_paragraph()
        r = p.add_run(label)
        r.bold = True
    for line in code.strip("\n").splitlines():
        p = doc.add_paragraph(style="CodeBlock")
        p.paragraph_format.left_indent = Inches(0.12)
        r = p.add_run(line if line else " ")
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")


def add_compare_table(doc, title, old_code, new_code, explanation):
    doc.add_heading(title, level=3)
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    headers = ["Code cũ", "Code mới"]
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_shading(cell, "E8EEF5")
        p = cell.paragraphs[0]
        p.style = doc.styles["TableText"]
        r = p.add_run(header)
        r.bold = True
    for i, code in enumerate([old_code, new_code]):
        cell = table.cell(1, i)
        for j, line in enumerate(code.strip("\n").splitlines()):
            p = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
            p.style = doc.styles["CodeBlock"]
            r = p.add_run(line if line else " ")
            r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    doc.add_paragraph(explanation)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_note(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, "BFC7D5")
    cell = table.cell(0, 0)
    set_cell_shading(cell, "EEF3FA")
    p = cell.paragraphs[0]
    r = p.add_run(title + ": ")
    r.bold = True
    p.add_run(text)


def add_image(doc, path, caption):
    doc.add_picture(str(path), width=Inches(6.7))
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    r = p.add_run(caption)
    r.italic = True
    r.font.size = Pt(8.5)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_shading(cell, "E8EEF5")
        p = cell.paragraphs[0]
        p.style = doc.styles["TableText"]
        r = p.add_run(header)
        r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            p.style = doc.styles["TableText"]
            p.add_run(value)
    doc.add_paragraph()


def build_doc():
    images = create_images()
    doc = setup_doc()
    doc.add_paragraph("Web Search Fix Guide", style="Title")
    doc.add_paragraph("Smart Study Planner - giải thích chi tiết lỗi search, code cũ, code mới và lý do sửa")

    doc.add_heading("1. Vấn đề cần hiểu", level=1)
    doc.add_paragraph(
        "Phần search trên web bị lỗi ở tầng frontend. Backend API search vẫn trả dữ liệu đúng khi gọi trực tiếp "
        "`/api/flashcards/search?keyword=join`. Vì vậy lỗi không nằm ở SQL Server hay FastAPI query chính, mà nằm ở cách giao diện web quản lý state, gọi search và render kết quả."
    )
    add_note(
        doc,
        "Kết luận ngắn",
        "Code cũ để search phụ thuộc vào việc gọi lại loadDashboard và request API search. Code mới tách rõ dữ liệu gốc và dữ liệu hiển thị: state.flashcards giữ toàn bộ flashcards, visibleFlashcards() lọc local theo keyword, renderSearchSummary() báo kết quả cho user.",
    )

    doc.add_heading("2. Triệu chứng khi nhìn từ người dùng", level=1)
    add_bullets(
        doc,
        [
            "Gõ keyword vào ô search nhưng cảm giác kết quả không rõ ràng hoặc không đổi như mong đợi.",
            "Không có dòng báo đang hiển thị bao nhiêu kết quả.",
            "Search phụ thuộc vào bấm nút hoặc Enter, chưa có phản hồi live khi đang gõ.",
            "Refresh/add/edit/delete có thể làm danh sách bị reload theo currentSearch cũ, khiến người dùng tưởng dữ liệu bị mất hoặc search bị kẹt.",
            "Trình duyệt có thể dùng cache app.js cũ, nên sau khi sửa code vẫn thấy hành vi cũ nếu không cache-bust.",
        ],
    )

    doc.add_heading("3. Code cũ chạy theo luồng nào?", level=1)
    add_image(doc, images["old_flow"], "Hình 1 - Luồng search cũ phụ thuộc API và reload toàn dashboard.")
    doc.add_paragraph(
        "Ở code cũ, khi user bấm Search, searchFlashcards() chỉ set state.currentSearch rồi gọi loadDashboard(). "
        "loadDashboard() quyết định URL: nếu có keyword thì gọi `/api/flashcards/search`, nếu không thì gọi `/api/flashcards`. "
        "Sau đó kết quả trả về được gán thẳng vào state.flashcards."
    )
    add_code(
        doc,
        """
async function loadDashboard() {
  setBusy(true);
  try {
    await checkHealth();
    const flashcardUrl = state.currentSearch
      ? `${api.flashcardSearch}?keyword=${encodeURIComponent(state.currentSearch)}`
      : api.flashcards;

    const [stats, flashcards, topics, dueReviews, weakTopics] = await Promise.all([
      requestJson(api.statistics),
      requestJson(flashcardUrl),
      requestJson(api.topics),
      requestJson(api.dueReviews),
      requestJson(api.weakTopics),
    ]);

    state.flashcards = flashcards;
    state.topics = topics;
    renderStats(stats);
    renderFlashcards(flashcards);
    renderTopics(topics);
    renderDueReviews(dueReviews);
    renderWeakTopics(weakTopics);
  } finally {
    setBusy(false);
  }
}
        """,
        "Đoạn quan trọng trong code cũ:",
    )

    doc.add_heading("4. Gốc lỗi là gì?", level=1)
    add_image(doc, images["root_cause"], "Hình 2 - Gốc lỗi: trộn dữ liệu gốc và dữ liệu đang hiển thị.")
    doc.add_paragraph(
        "Một app frontend thường cần phân biệt hai thứ: dữ liệu gốc và dữ liệu đang hiển thị. "
        "Dữ liệu gốc là toàn bộ flashcards đã lấy từ server. Dữ liệu đang hiển thị có thể là toàn bộ, hoặc chỉ là kết quả search/filter. "
        "Code cũ dùng state.flashcards cho cả hai vai trò. Khi search, state.flashcards có thể trở thành danh sách đã lọc; khi clear/refresh lại gọi API; khi edit/delete lại tìm trong state.flashcards hiện tại. Cách này vẫn có thể chạy, nhưng dễ sinh hành vi khó hiểu."
    )
    add_table(
        doc,
        ["Vấn đề", "Trong code cũ", "Hậu quả"],
        [
            ("Source of truth không rõ", "state.flashcards lúc là all data, lúc là search results.", "Các action như edit/delete/refresh phụ thuộc trạng thái hiện tại, dễ khó debug."),
            ("Search gọi lại toàn dashboard", "searchFlashcards -> loadDashboard -> Promise.all nhiều API.", "Mỗi lần search kéo theo statistics/topics/reviews/weak topics, không cần thiết."),
            ("Không có summary", "HTML không có searchSummary.", "User không biết đang xem toàn bộ hay kết quả lọc."),
            ("Không có live search", "Chỉ search khi click hoặc Enter.", "Gõ xong nhưng chưa bấm có cảm giác không hoạt động."),
            ("Có thể bị cache JS", "index.html dùng /static/app.js không version.", "Browser có thể giữ file JS cũ sau khi sửa."),
        ],
    )

    doc.add_heading("5. Vì sao API search đúng mà web vẫn bị lỗi?", level=1)
    doc.add_paragraph(
        "API search là backend. Frontend search là trải nghiệm trên trình duyệt. Một endpoint đúng chưa đảm bảo UI đúng. "
        "Trong lần kiểm tra, gọi trực tiếp endpoint `/api/flashcards/search?keyword=join` trả về 3 flashcards, nghĩa là SQL và FastAPI hoạt động. "
        "Lỗi nằm ở cách UI gọi API, lưu state, render kết quả và phản hồi cho user."
    )
    add_note(
        doc,
        "Bài học quan trọng",
        "Khi debug web app, luôn tách câu hỏi: API có trả đúng không? Frontend có gọi đúng không? Frontend có render đúng không? Browser có đang dùng file JS mới không?",
    )

    doc.add_heading("6. Code mới sửa theo hướng nào?", level=1)
    add_image(doc, images["new_flow"], "Hình 3 - Luồng search mới: filter local, dữ liệu gốc ổn định.")
    doc.add_paragraph(
        "Code mới đổi tư duy: loadDashboard luôn lấy toàn bộ flashcards từ `/api/flashcards`, lưu vào state.flashcards như nguồn dữ liệu gốc. "
        "Khi search, app không gọi lại server nữa mà lọc local bằng visibleFlashcards(). Vì số lượng flashcard của app học tập cá nhân còn nhỏ, cách này nhanh, dễ hiểu và ít lỗi state hơn."
    )

    add_compare_table(
        doc,
        "6.1. Thêm searchSummary vào DOM",
        """
const elements = {
  searchInput: document.querySelector("#searchInput"),
  searchButton: document.querySelector("#searchButton"),
  clearSearchButton: document.querySelector("#clearSearchButton"),
  flashcardList: document.querySelector("#flashcardList"),
};
        """,
        """
const elements = {
  searchInput: document.querySelector("#searchInput"),
  searchButton: document.querySelector("#searchButton"),
  clearSearchButton: document.querySelector("#clearSearchButton"),
  searchSummary: document.querySelector("#searchSummary"),
  flashcardList: document.querySelector("#flashcardList"),
};
        """,
        "Code mới thêm searchSummary để JavaScript có thể cập nhật dòng trạng thái: đang hiển thị toàn bộ hay bao nhiêu kết quả theo keyword.",
    )

    add_compare_table(
        doc,
        "6.2. Thêm vùng hiển thị số kết quả trong HTML",
        """
<div class="section-heading">
  ...
</div>
<div id="flashcardList" class="card-list"></div>
        """,
        """
<div class="section-heading">
  ...
</div>
<p id="searchSummary" class="search-summary">
  Showing all flashcards.
</p>
<div id="flashcardList" class="card-list"></div>
        """,
        "Trước đây user không có feedback. Sau khi sửa, user biết search đã ăn hay chưa qua dòng `Showing 3 results for \"join\"`.",
    )

    add_compare_table(
        doc,
        "6.3. loadDashboard cũ gọi API search, code mới luôn lấy all flashcards",
        """
const flashcardUrl = state.currentSearch
  ? `${api.flashcardSearch}?keyword=${encodeURIComponent(state.currentSearch)}`
  : api.flashcards;

const [stats, flashcards, topics, dueReviews, weakTopics] = await Promise.all([
  requestJson(api.statistics),
  requestJson(flashcardUrl),
  requestJson(api.topics),
  requestJson(api.dueReviews),
  requestJson(api.weakTopics),
]);

state.flashcards = flashcards;
renderFlashcards(flashcards);
        """,
        """
const [stats, flashcards, topics, dueReviews, weakTopics] = await Promise.all([
  requestJson(api.statistics),
  requestJson(api.flashcards),
  requestJson(api.topics),
  requestJson(api.dueReviews),
  requestJson(api.weakTopics),
]);

state.flashcards = flashcards;
const filteredFlashcards = visibleFlashcards();
renderFlashcards(filteredFlashcards);
renderSearchSummary(filteredFlashcards.length);
        """,
        "Điểm khác biệt lớn nhất: state.flashcards giờ luôn là danh sách gốc. Danh sách hiển thị được tính riêng bằng visibleFlashcards(). Đây là cách tách source data và view data.",
    )

    doc.add_heading("7. Ba function mới quan trọng", level=1)
    add_code(
        doc,
        """
function normalizeSearchText(value) {
  return String(value ?? "")
    .normalize("NFD")
    .replace(/[\\u0300-\\u036f]/g, "")
    .toLowerCase()
    .trim();
}
        """,
        "normalizeSearchText:",
    )
    doc.add_paragraph(
        "Function này chuẩn hóa text trước khi so sánh: chuyển null/undefined thành chuỗi rỗng, tách dấu Unicode bằng normalize('NFD'), bỏ dấu tiếng Việt cơ bản, chuyển về chữ thường và bỏ khoảng trắng đầu/cuối. "
        "Nhờ vậy search `join`, `JOIN`, hoặc keyword có dấu/không dấu sẽ ổn định hơn."
    )

    add_code(
        doc,
        """
function visibleFlashcards() {
  const keyword = normalizeSearchText(state.currentSearch);
  if (!keyword) return state.flashcards;

  return state.flashcards.filter((card) => {
    const searchableText = normalizeSearchText(
      [card.question, card.answer, card.topicName, card.subjectName].join(" ")
    );
    return searchableText.includes(keyword);
  });
}
        """,
        "visibleFlashcards:",
    )
    doc.add_paragraph(
        "Function này tạo danh sách flashcard đang hiển thị. Nếu không có keyword thì trả toàn bộ state.flashcards. "
        "Nếu có keyword, nó gộp question, answer, topicName, subjectName thành một chuỗi lớn rồi kiểm tra chuỗi đó có chứa keyword không."
    )

    add_code(
        doc,
        """
function renderSearchSummary(visibleCount) {
  if (!state.currentSearch) {
    elements.searchSummary.textContent = `Showing all ${state.flashcards.length} flashcards.`;
    return;
  }

  elements.searchSummary.textContent =
    `Showing ${visibleCount} result${visibleCount === 1 ? "" : "s"} for "${state.currentSearch}".`;
}
        """,
        "renderSearchSummary:",
    )
    doc.add_paragraph(
        "Function này không ảnh hưởng dữ liệu, nhưng rất quan trọng cho UX. Nó giúp user biết search đang ở trạng thái nào. "
        "Khi debug frontend, dòng trạng thái kiểu này cũng giúp lập trình viên nhìn ra app có đang nhận keyword không."
    )

    doc.add_heading("8. Search và Clear thay đổi thế nào?", level=1)
    add_compare_table(
        doc,
        "8.1. searchFlashcards",
        """
function searchFlashcards() {
  state.currentSearch = elements.searchInput.value.trim();
  loadDashboard();
}
        """,
        """
function searchFlashcards() {
  state.currentSearch = elements.searchInput.value.trim();
  const filteredFlashcards = visibleFlashcards();
  renderFlashcards(filteredFlashcards);
  renderSearchSummary(filteredFlashcards.length);
}
        """,
        "Code cũ search bằng cách reload toàn dashboard. Code mới chỉ filter và render lại danh sách flashcards. Nhanh hơn, ít API thừa hơn, ít state rối hơn.",
    )
    add_compare_table(
        doc,
        "8.2. clearSearch",
        """
function clearSearch() {
  state.currentSearch = "";
  elements.searchInput.value = "";
  loadDashboard();
}
        """,
        """
function clearSearch() {
  state.currentSearch = "";
  elements.searchInput.value = "";
  renderFlashcards(state.flashcards);
  renderSearchSummary(state.flashcards.length);
}
        """,
        "Clear không cần gọi server. Vì state.flashcards đã giữ toàn bộ dữ liệu gốc, chỉ cần render lại danh sách gốc.",
    )

    doc.add_heading("9. Vì sao thêm input event?", level=1)
    add_compare_table(
        doc,
        "9.1. Event search",
        """
elements.searchButton.addEventListener("click", searchFlashcards);
elements.searchInput.addEventListener("keydown", (event) => {
  if (event.key === "Enter") searchFlashcards();
});
        """,
        """
elements.searchButton.addEventListener("click", searchFlashcards);
elements.searchInput.addEventListener("keydown", (event) => {
  if (event.key === "Enter") searchFlashcards();
});
elements.searchInput.addEventListener("input", () => {
  state.currentSearch = elements.searchInput.value.trim();
  const filteredFlashcards = visibleFlashcards();
  renderFlashcards(filteredFlashcards);
  renderSearchSummary(filteredFlashcards.length);
});
        """,
        "Event input giúp search chạy ngay khi user gõ. Đây là cải thiện UX: user không cần đoán là phải bấm Search hay Enter.",
    )

    doc.add_heading("10. Vì sao thêm cache-bust vào app.js và styles.css?", level=1)
    add_compare_table(
        doc,
        "10.1. Cache-bust static assets",
        """
<link rel="stylesheet" href="/static/styles.css" />
...
<script src="/static/app.js"></script>
        """,
        """
<link rel="stylesheet" href="/static/styles.css?v=search-fix" />
...
<script src="/static/app.js?v=search-fix"></script>
        """,
        "Trình duyệt có thể cache file app.js cũ. Khi thêm `?v=search-fix`, URL file thay đổi, browser buộc tải bản mới. Đây là mẹo phổ biến khi sửa frontend mà người dùng vẫn thấy hành vi cũ.",
    )

    doc.add_heading("11. Minh họa UI sau khi sửa", level=1)
    add_image(doc, images["mock_ui"], "Hình 4 - Minh họa trạng thái search sau fix: keyword `join` trả 3 kết quả và có dòng summary.")

    dashboard = ROOT / "docs" / "assets" / "web_dashboard.png"
    if dashboard.exists():
        add_image(doc, dashboard, "Hình 5 - Screenshot web dashboard hiện tại sau khi thêm search summary.")

    doc.add_heading("12. So sánh tư duy cũ và mới", level=1)
    add_table(
        doc,
        ["Khía cạnh", "Code cũ", "Code mới"],
        [
            ("Nguồn dữ liệu", "state.flashcards có thể là all hoặc filtered.", "state.flashcards luôn là all data."),
            ("Dữ liệu hiển thị", "Dùng luôn flashcards trả về từ API.", "Tạo riêng bằng visibleFlashcards()."),
            ("Search", "Gọi API search và reload dashboard.", "Filter local ngay trên frontend."),
            ("Clear", "Gọi lại loadDashboard.", "Render lại state.flashcards."),
            ("Phản hồi user", "Không có summary.", "Có Showing all / Showing N results."),
            ("Typing", "Chỉ click/Enter.", "Có live filter khi input thay đổi."),
            ("Cache", "app.js URL cố định.", "Có ?v=search-fix để tải JS/CSS mới."),
        ],
    )

    doc.add_heading("13. Khi nào nên search ở frontend, khi nào nên search ở backend?", level=1)
    add_table(
        doc,
        ["Trường hợp", "Nên dùng", "Lý do"],
        [
            ("Dữ liệu ít, app cá nhân, vài chục/vài trăm flashcard", "Frontend filter", "Nhanh, đơn giản, không gọi API liên tục."),
            ("Dữ liệu lớn, hàng chục nghìn bản ghi", "Backend search", "Không nên tải hết data lên browser."),
            ("Cần phân trang, phân quyền, lọc theo user", "Backend search", "Database/server phải quyết định dữ liệu nào user được xem."),
            ("Cần trải nghiệm tức thì khi gõ", "Frontend filter hoặc debounce backend", "Tránh gọi server quá nhiều."),
        ],
    )
    doc.add_paragraph(
        "Với dự án hiện tại, frontend filter là hợp lý vì đây là app học tập cá nhân, dữ liệu nhỏ và cần dễ hiểu. "
        "Nếu sau này có nhiều user hoặc dữ liệu lớn, có thể quay lại backend search nhưng vẫn phải giữ tư duy tách `all data`, `search result`, `currentSearch` rõ ràng."
    )

    doc.add_heading("14. Bài học rút ra để tự debug lần sau", level=1)
    add_numbered(
        doc,
        [
            "Kiểm tra API trước: gọi endpoint trực tiếp xem backend có trả đúng không.",
            "Nếu API đúng, kiểm tra frontend có gọi đúng URL không.",
            "Kiểm tra state: biến nào là dữ liệu gốc, biến nào là dữ liệu hiển thị.",
            "Kiểm tra render: sau khi state đổi, UI có gọi render lại không.",
            "Kiểm tra UX: user có biết search đang có bao nhiêu kết quả không.",
            "Kiểm tra cache: browser có đang dùng app.js cũ không.",
            "Sau khi sửa, test keyword có kết quả và keyword không có kết quả.",
        ],
    )

    doc.add_heading("15. Checklist kiểm tra search sau fix", level=1)
    add_bullets(
        doc,
        [
            "Mở http://127.0.0.1:8000.",
            "Ô search ban đầu có dòng `Showing all 10 flashcards.`.",
            "Gõ `join` thì danh sách còn các card liên quan SQL Joins hoặc INNER JOIN.",
            "Dòng summary đổi thành `Showing 3 results for \"join\"` nếu database đang giống lúc test.",
            "Gõ keyword không tồn tại thì hiện `No flashcards found.`.",
            "Bấm Clear thì hiện lại toàn bộ flashcards.",
            "Sau khi add/edit/delete, nếu search còn keyword thì list vẫn lọc theo keyword đó.",
        ],
    )

    doc.add_heading("16. Tóm tắt thật ngắn", level=1)
    doc.add_paragraph(
        "Search cũ bị yếu vì trộn dữ liệu gốc và kết quả search trong cùng state.flashcards, đồng thời search phải reload cả dashboard qua API. "
        "Search mới giữ state.flashcards là dữ liệu gốc, dùng visibleFlashcards() để lọc danh sách hiển thị, thêm search summary và input event để user thấy kết quả ngay. "
        "Đây là bài học quan trọng về state management trong frontend."
    )

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_doc()
    print(OUTPUT)
